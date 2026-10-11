from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import shutil
import tempfile
import tomllib
import xml.etree.ElementTree as ET
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from zipfile import ZipFile

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from hangbo.precheck.phase_column import preprocess_phase_column_if_needed
from hangbo.summary.table import (
    SUMMARY_COLUMNS,
    ReportSnapshot,
    SummaryRowResult,
    build_summary_rows,
    build_total_values,
    style_summary_worksheet,
)
from hangbo.survey.customer_category_rules import (
    CUSTOMER_CATEGORY_RULE_BY_NAME,
    DISPLAY_ORDERED_CUSTOMER_CATEGORY_RULES,
)
from hangbo.survey.stats import (
    SurveyStatistics,
    build_role_definition_from_customer_category_rule,
    compute_role_stats,
    merge_survey_statistics,
    normalize_calculation_mode,
    validate_dataframe,
)

SECTION_COLUMNS = {
    "会展服务": "产品服务", "会场服务": "产品服务", "入住服务": "产品服务",
    "旅游服务": "产品服务", "硬件设施": "硬件设施", "配套服务": "配套服务",
    "智慧场馆": "智慧场馆/服务", "智慧服务": "智慧场馆/服务", "餐饮服务": "餐饮服务",
}


def mean(values: list[float]) -> float:
    return math.fsum(values) / len(values)


def period_labels_for_batch(batch: str) -> tuple[str, ...]:
    batch = batch.strip().replace("–", "-").replace("—", "-")
    quarter = re.fullmatch(r"Q([1-4])", batch.upper())
    if quarter:
        start, end = 3 * int(quarter[1]) - 2, 3 * int(quarter[1])
    else:
        match = re.fullmatch(r"(\d{1,2})(?:-(\d{1,2}))?月", batch)
        if match is None:
            raise ValueError("批次须为月份、月份范围或 Q1-Q4；自定义批次请使用 --periods。")
        start, end = int(match[1]), int(match[2] or match[1])
    if not 1 <= start <= end <= 12:
        raise ValueError("月份范围须在 1-12 月内，且起始月份不大于结束月份。")
    if start <= 2:
        if start != 1 or end < 2:
            raise ValueError("默认 1-2 月为一个完整周期；独立 1 月或 2 月请使用 --periods。")
        return ("1-2月", *(f"{month}月" for month in range(3, end + 1)))
    return tuple(f"{month}月" for month in range(start, end + 1))


def validate_period_labels(periods: tuple[str, ...]) -> tuple[str, ...]:
    normalized: list[str] = []
    covered: set[int] = set()
    for period in periods:
        label = period.strip().replace("–", "-").replace("—", "-")
        if label == "1-2月":
            months = {1, 2}
        else:
            match = re.fullmatch(r"(0?[1-9]|1[0-2])月", label)
            if match is None:
                raise ValueError(f"不是独立计算周期：{period}；仅支持单月或 1-2月。")
            month = int(match[1])
            label, months = f"{month}月", {month}
        if covered & months:
            raise ValueError(f"重复或重叠的计算周期：{period}")
        covered.update(months)
        normalized.append(label)
    if not normalized:
        raise ValueError("至少选择一个计算周期。")
    return tuple(normalized)


@dataclass(frozen=True)
class RowAggregation:
    rows: tuple[SummaryRowResult, ...]
    total_values: dict[str, float | None]
    row_weights: dict[tuple[str, str], float]
    period_totals: dict[str, dict[str, float | None]]
    rounding_differences: dict[str, float]


def aggregate_period_summary_rows(
    period_rows: dict[str, tuple[SummaryRowResult, ...]],
) -> RowAggregation:
    if not period_rows:
        raise ValueError("至少需要一个周期。")
    definitions = next(iter(period_rows.values()))
    names = [(row.category_label, row.display_name) for row in definitions]
    if any([(row.category_label, row.display_name) for row in rows] != names for rows in period_rows.values()):
        raise ValueError("各期客户类型结构不一致。")
    totals = {period: build_total_values(rows) for period, rows in period_rows.items()}
    if any(values["总分"] is None for values in totals.values()):
        raise ValueError("所选周期没有有效满意度；不能静默跳过整个周期。")
    weights: dict[tuple[str, str], float] = {}
    contributions: dict[tuple[str, str], float] = {}
    total_values: dict[str, float | None] = {}
    differences: dict[str, float] = {}
    for column in SUMMARY_COLUMNS:
        valid_periods = [rows for rows in period_rows.values() if any(row.display_name and row.values[column] is not None for row in rows)]
        total_values[column] = mean([values[column] for values in totals.values() if values[column] is not None]) if valid_periods else None
        for rows in valid_periods:
            valid_rows = [row for row in rows if row.display_name and row.values[column] is not None]
            weight = 1 / len(valid_periods) / len(valid_rows)
            for row in valid_rows:
                key = row.display_name, column
                weights[key] = weights.get(key, 0) + weight
                contributions[key] = contributions.get(key, 0) + weight * row.values[column]
        differences[column] = (total_values[column] or 0) - math.fsum(value for (name, col), value in contributions.items() if col == column)
    rows = tuple(
        SummaryRowResult(
            category_label=row.category_label, display_name=row.display_name,
            values={column: contributions[(row.display_name, column)] / weights[(row.display_name, column)] if weights.get((row.display_name, column)) else None for column in SUMMARY_COLUMNS},
            applicable_columns=row.applicable_columns,
        )
        for row in definitions
    )
    return RowAggregation(rows, total_values, weights, totals, differences)


@dataclass(frozen=True)
class PeriodCalculation:
    label: str
    rows: tuple[SummaryRowResult, ...]
    components: dict[str, tuple[SurveyStatistics, ...]]
    source_hashes: dict[str, str]
    notices: tuple[str, ...]
    saved_total: float | None


def read_saved_total(path: Path) -> float | None:
    if not path.exists():
        return None
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for row in workbook.active.iter_rows(values_only=True):
            if row[0] == "总分" or row[1] == "总分":
                return float(row[2]) if isinstance(row[2], (int, float)) else None
    finally:
        workbook.close()
    return None


def calculate_period(
    raw_dir: Path, label: str, *, sheet_name: str = "问卷数据", calculation_mode: str = "template",
    saved_summary_path: Path | None = None,
) -> PeriodCalculation:
    if not raw_dir.is_dir():
        raise ValueError(f"缺少所选周期原始数据目录：{raw_dir}")
    reports: list[ReportSnapshot] = []
    components: dict[str, tuple[SurveyStatistics, ...]] = {}
    hashes: dict[str, str] = {}
    notices: list[str] = []
    frames: dict[str, pd.DataFrame] = {}
    with tempfile.TemporaryDirectory(prefix="hangbo-multi-period-") as temporary:
        for filename in dict.fromkeys(rule.source_file_name for rule in DISPLAY_ORDERED_CUSTOMER_CATEGORY_RULES):
            source = raw_dir / filename
            if not source.is_file():
                notices.append(f"缺少来源文件：{filename}，相应客户类型无数据。")
                continue
            hashes[filename] = hashlib.sha256(source.read_bytes()).hexdigest()
            copied = Path(temporary) / filename
            shutil.copy2(source, copied)
            notice = preprocess_phase_column_if_needed(copied, sheet_name)
            if notice:
                notices.append(f"{filename}：第三列 X 期标记已在临时副本上移至末列。")
            frames[filename] = pd.read_excel(copied, sheet_name=sheet_name)
            if hashes[filename] != hashlib.sha256(source.read_bytes()).hexdigest():
                raise ValueError(f"计算期间来源文件发生变化：{source}")
        for rule in DISPLAY_ORDERED_CUSTOMER_CATEGORY_RULES:
            if rule.source_file_name not in frames:
                continue
            frame = frames[rule.source_file_name]
            rules = [CUSTOMER_CATEGORY_RULE_BY_NAME[name] for name in rule.aggregate_rule_names] if rule.is_aggregate else [rule]
            calculated: list[SurveyStatistics] = []
            for component in rules:
                definition = build_role_definition_from_customer_category_rule(component)
                validate_dataframe(frame, definition)
                calculated.append(compute_role_stats(frame, definition, calculation_mode=calculation_mode))
            stats = merge_survey_statistics(rule.name, calculated) if rule.is_aggregate else calculated[0]
            components[rule.customer_category] = tuple(calculated)
            if stats.matched_row_count == 0:
                notices.append(f"{rule.customer_category}：无样本。")
            elif stats.satisfaction is None:
                notices.append(f"{rule.customer_category}：有样本，但无有效满意度回答。")
            reports.append(ReportSnapshot(
                role_name=rule.name, source_path=(raw_dir / f"{rule.name}.xlsx").resolve(),
                total=stats.satisfaction,
                sections={section.name: section.satisfaction for section in stats.sections},
                metrics={metric.name: metric.satisfaction for section in stats.sections for metric in section.metrics},
            ))
    rows = build_summary_rows(tuple(reports))
    if build_total_values(rows)["总分"] is None:
        raise ValueError(f"周期 {label} 没有有效满意度，已停止汇总。")
    saved = read_saved_total(saved_summary_path) if saved_summary_path else None
    current = build_total_values(rows)["总分"]
    if saved is not None and not math.isclose(saved, current, abs_tol=1e-12):
        notices.append(f"历史表总分 {saved:.2f} 与本次原始数据重算 {current:.2f} 不同；本次统一使用重算值。")
    return PeriodCalculation(label, rows, components, hashes, tuple(notices), saved)


@dataclass(frozen=True)
class SectionTerm:
    period: str
    category: str
    customer_type: str
    component: str
    section: str
    column: str
    score: float
    period_weight: float
    type_weight: float
    component_weight: float
    section_weight: float

    @property
    def weight(self) -> float:
        return self.period_weight * self.type_weight * self.component_weight * self.section_weight

    @property
    def display_weight(self) -> float:
        return self.weight if self.section_weight else self.period_weight * self.type_weight


@dataclass(frozen=True)
class TypeTerm:
    period: str
    customer_type: str
    score: float
    reconstructed: float
    weight: float

    @property
    def rounding(self) -> float:
        return self.score - self.reconstructed


@dataclass(frozen=True)
class MultiPeriodResult:
    periods: tuple[PeriodCalculation, ...]
    rows: tuple[SummaryRowResult, ...]
    aggregation: RowAggregation
    sections: tuple[SectionTerm, ...]
    types: tuple[TypeTerm, ...]
    cell_weights: dict[tuple[str, str], float]
    column_scores: dict[str, float | None]
    checks: dict[str, float]


def build_multi_period_result(periods: tuple[PeriodCalculation, ...]) -> MultiPeriodResult:
    if len({period.label for period in periods}) != len(periods):
        raise ValueError("计算周期重复。")
    aggregation = aggregate_period_summary_rows({period.label: period.rows for period in periods})
    sections: list[SectionTerm] = []
    types: list[TypeTerm] = []
    for period in periods:
        valid_rows = [row for row in period.rows if row.display_name and row.values["总分"] is not None]
        for row in valid_rows:
            type_sections: list[SectionTerm] = []
            components = [stats for stats in period.components[row.display_name] if stats.satisfaction is not None]
            if not components:
                raise ValueError(f"缺少 {period.label}/{row.display_name} 的计分组件。")
            for component in components:
                valid_sections = [section for section in component.sections if section.satisfaction is not None]
                if not valid_sections:
                    raise ValueError(f"缺少 {component.role_name} 的有效一级指标。")
                for section in valid_sections:
                    if section.name not in SECTION_COLUMNS:
                        raise ValueError(f"一级指标未映射到互斥展示列：{section.name}")
                    type_sections.append(SectionTerm(
                        period.label, row.category_label, row.display_name, component.role_name,
                        section.name, SECTION_COLUMNS[section.name], section.satisfaction,
                        1 / len(periods), 1 / len(valid_rows), 1 / len(components), 1 / len(valid_sections),
                    ))
            weight = 1 / len(periods) / len(valid_rows)
            reconstructed = math.fsum(term.score * term.weight for term in type_sections) / weight
            types.append(TypeTerm(period.label, row.display_name, row.values["总分"], reconstructed, weight))
            sections.extend(type_sections)
            if row.values["餐饮服务"] is not None and not any(term.column == "餐饮服务" for term in type_sections):
                sections.append(SectionTerm(
                    period.label, row.category_label, row.display_name, "配套服务子项", "其中：餐饮服务",
                    "餐饮服务", row.values["餐饮服务"], 1 / len(periods), 1 / len(valid_rows), 1, 0,
                ))
    cell_weights: dict[tuple[str, str], float] = {}
    rows: list[SummaryRowResult] = []
    column_scores: dict[str, float | None] = {}
    for row in aggregation.rows:
        values = {"总分": row.values["总分"]}
        for column in SUMMARY_COLUMNS[1:]:
            matching = [term for term in sections if term.customer_type == row.display_name and term.column == column]
            cell_weights[row.display_name, column] = math.fsum(term.weight for term in matching)
            weight = math.fsum(term.display_weight for term in matching)
            values[column] = math.fsum(term.score * term.display_weight for term in matching) / weight if weight else None
        rows.append(SummaryRowResult(row.category_label, row.display_name, values, row.applicable_columns))
    for column in SUMMARY_COLUMNS[1:]:
        matching = [term for term in sections if term.column == column and term.weight]
        weight = math.fsum(term.weight for term in matching)
        column_scores[column] = math.fsum(term.score * term.weight for term in matching) / weight if weight else None
    target = aggregation.total_values["总分"]
    contributions = math.fsum(term.score * term.weight for term in sections)
    rounding = math.fsum(term.rounding * term.weight for term in types)
    checks = {
        "客户类型权重和差额": math.fsum(term.weight for term in types) - 1,
        "一级指标权重和差额": math.fsum(term.weight for term in sections) - 1,
        "纵向总分差额": math.fsum(term.score * term.weight for term in types) + aggregation.rounding_differences["总分"] - target,
        "横向总分差额": contributions + rounding + aggregation.rounding_differences["总分"] - target,
    }
    for row in rows:
        weight = aggregation.row_weights.get((row.display_name, "总分"), 0)
        if weight:
            section_sum = math.fsum((row.values[column] or 0) * cell_weights[row.display_name, column] for column in SUMMARY_COLUMNS[1:])
            delta = math.fsum(term.rounding * term.weight for term in types if term.customer_type == row.display_name)
            checks[f"{row.display_name}横向差额"] = (section_sum + delta) / weight - row.values["总分"]
    if any(abs(value) > 1e-10 for value in checks.values()):
        raise ValueError(f"多周期横纵核对失败：{checks}")
    return MultiPeriodResult(periods, tuple(rows), aggregation, tuple(sections), tuple(types), cell_weights, column_scores, checks)


class FormulaWriter:
    def __init__(self) -> None:
        self.cached: dict[tuple[str, str], float | str] = {}

    def set(self, sheet, coordinate: str, formula: str, value: float | str) -> None:
        sheet[coordinate] = formula
        self.cached[sheet.title, coordinate] = value

    def save(self, workbook: Workbook, path: Path) -> None:
        workbook.save(path)
        # OOXML supports a formula and its cached result together; read-only consumers need both.
        namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        ET.register_namespace("", namespace)
        with ZipFile(path) as archive:
            members = [(info, archive.read(info.filename)) for info in archive.infolist()]
        sheet_paths = {f"xl/worksheets/sheet{index}.xml": sheet.title for index, sheet in enumerate(workbook, 1)}
        with ZipFile(path, "w") as archive:
            for info, contents in members:
                if info.filename in sheet_paths:
                    root = ET.fromstring(contents)
                    for cell in root.iter(f"{{{namespace}}}c"):
                        key = sheet_paths[info.filename], cell.attrib["r"]
                        if key not in self.cached:
                            continue
                        value = self.cached[key]
                        cached = cell.find(f"{{{namespace}}}v")
                        if cached is None:
                            cached = ET.SubElement(cell, f"{{{namespace}}}v")
                        cached.text = str(value)
                        if isinstance(value, str):
                            cell.set("t", "str")
                    contents = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                archive.writestr(info, contents)


def style_calculation_sheet(sheet) -> None:
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for row in sheet:
        for cell in row:
            cell.font = Font(name="Arial", size=11)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor="B32046")
                cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            elif isinstance(cell.value, (float, int)) or cell.data_type == "f":
                cell.number_format = "0.000000"
    sheet.row_dimensions[1].height = 38
    for column in range(1, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 19


def write_multi_period_workbook(result: MultiPeriodResult, path: Path, *, year: str, batch: str) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "汇总表"
    cycle = workbook.create_sheet("周期计算")
    detail = workbook.create_sheet("权重计算")
    formulas = FormulaWriter()
    cycle.append(["周期", "数据来源", *SUMMARY_COLUMNS, "有效客户类型数", "类型总分均值", "周期舍入差额", "周期权重", "周期总分贡献", "周期舍入贡献", "历史保存总分", "重算与历史差异"])
    detail.append(["周期", "客户大类", "样本类型", "来源组件", "一级指标或子项", "展示列", "单期得分", "周期权重", "期内类型权重", "组件权重", "一级指标权重", "累计计分权重", "得分贡献", "展示权重", "展示贡献", None, "周期", "样本类型", "单期类型总分", "一级指标重建总分", "类型舍入差额", "累计类型权重", "类型总分贡献", "类型舍入贡献"])
    for index, term in enumerate(result.sections, 2):
        detail.append([term.period, term.category, term.customer_type, term.component, term.section, term.column, term.score, term.period_weight, term.type_weight, term.component_weight, term.section_weight])
        formulas.set(detail, f"L{index}", f"=PRODUCT(H{index}:K{index})", term.weight)
        formulas.set(detail, f"M{index}", f"=G{index}*L{index}", term.score * term.weight)
        formulas.set(detail, f"N{index}", f"=IF(K{index}=0,H{index}*I{index},L{index})", term.display_weight)
        formulas.set(detail, f"O{index}", f"=G{index}*N{index}", term.score * term.display_weight)
    end = len(result.sections) + 1
    type_end = len(result.types) + 1
    ref = lambda col: f"'权重计算'!${col}$2:${col}${end}"
    tref = lambda col: f"'权重计算'!${col}$2:${col}${type_end}"
    for index, term in enumerate(result.types, 2):
        for column, value in zip(("Q", "R", "S", "V"), (term.period, term.customer_type, term.score, term.weight)):
            detail[f"{column}{index}"] = value
        formulas.set(detail, f"T{index}", f"=SUMIFS($M$2:$M${end},$A$2:$A${end},Q{index},$C$2:$C${end},R{index})/V{index}", term.reconstructed)
        formulas.set(detail, f"U{index}", f"=S{index}-T{index}", term.rounding)
        formulas.set(detail, f"W{index}", f"=S{index}*V{index}", term.score * term.weight)
        formulas.set(detail, f"X{index}", f"=U{index}*V{index}", term.rounding * term.weight)
    for index, period in enumerate(result.periods, 2):
        totals = result.aggregation.period_totals[period.label]
        cycle.append([period.label, "原始数据重算", *(totals[column] for column in SUMMARY_COLUMNS)])
        valid = [row.values["总分"] for row in period.rows if row.display_name and row.values["总分"] is not None]
        unrounded = mean(valid)
        cycle[f"I{index}"] = len(valid)
        formulas.set(cycle, f"C{index}", f"=ROUND(J{index},2)", totals["总分"])
        formulas.set(cycle, f"J{index}", f"=SUMIFS({tref('W')},{tref('Q')},A{index})/L{index}", unrounded)
        formulas.set(cycle, f"K{index}", f"=C{index}-J{index}", totals["总分"] - unrounded)
        cycle[f"L{index}"] = 1 / len(result.periods)
        formulas.set(cycle, f"M{index}", f"=C{index}*L{index}", totals["总分"] / len(result.periods))
        formulas.set(cycle, f"N{index}", f"=K{index}*L{index}", (totals["总分"] - unrounded) / len(result.periods))
        if period.saved_total is not None:
            cycle[f"O{index}"] = period.saved_total
            formulas.set(cycle, f"P{index}", f"=C{index}-O{index}", totals["总分"] - period.saved_total)
    period_end = len(result.periods) + 1
    trend_row = period_end + 1
    cycle[f"A{trend_row}"] = "各期列总分均值"
    for col, name in enumerate(SUMMARY_COLUMNS, 3):
        if result.aggregation.total_values[name] is not None:
            letter = get_column_letter(col)
            formulas.set(cycle, f"{letter}{trend_row}", f"=AVERAGE({letter}2:{letter}{period_end})", result.aggregation.total_values[name])

    style_summary_worksheet(summary, result.rows)
    summary.unmerge_cells("A1:H1")
    summary.merge_cells("A1:K1")
    summary["A1"] = f"{year}年{batch}客户类型满意度多周期汇总表"
    summary["I2"], summary["J2"], summary["K2"] = "类型权重", "舍入差额", "餐饮计分口径"
    summary["J2"].comment = Comment("类型行列示该类型的舍入差额；总分行列示类型舍入贡献与周期舍入贡献之和。", "计算口径")
    total_row = len(result.rows) + 3
    weights_header = total_row + 6
    summary[f"A{weights_header - 1}"] = "计分权重与横向核对"
    for column, name in enumerate(("客户大类", "样本类型", "类型权重", *SUMMARY_COLUMNS[1:], "横向核对差额"), 1):
        summary.cell(weights_header, column, name)
    for index, row in enumerate(result.rows, 3):
        if not row.display_name:
            continue
        weight_row = weights_header + index - 2
        summary[f"A{weight_row}"] = row.category_label
        summary[f"B{weight_row}"] = row.display_name
        weight = result.aggregation.row_weights.get((row.display_name, "总分"), 0)
        formulas.set(summary, f"I{index}", f"=SUMIFS({tref('V')},{tref('R')},$B{index})", weight)
        formulas.set(summary, f"C{weight_row}", f"=I{index}", weight)
        delta = math.fsum(term.rounding * term.weight for term in result.types if term.customer_type == row.display_name) / weight if weight else 0
        formulas.set(summary, f"J{index}", f'=IF(I{index}=0,"",SUMIFS({tref("X")},{tref("R")},$B{index})/I{index})', delta if weight else "")
        formulas.set(summary, f"C{index}", f'=IF(I{index}=0,"",SUMIFS({tref("W")},{tref("R")},$B{index})/I{index})', row.values["总分"] if weight else "")
        for col, name in enumerate(SUMMARY_COLUMNS[1:], 4):
            letter = get_column_letter(col)
            cell_weight = result.cell_weights[row.display_name, name]
            condition = f'{ref("C")},$B{index},{ref("F")},{letter}$2'
            formulas.set(summary, f"{letter}{weight_row}", f"=SUMIFS({ref('L')},{condition})", cell_weight)
            if row.values[name] is not None:
                formulas.set(summary, f"{letter}{index}", f"=SUMIFS({ref('O')},{condition})/SUMIFS({ref('N')},{condition})", row.values[name])
        formulas.set(summary, f"I{weight_row}", f'=IF(I{index}=0,"",C{index}-SUMPRODUCT(D{index}:H{index},D{weight_row}:H{weight_row})/I{index}-J{index})', result.checks.get(f"{row.display_name}横向差额", ""))
        summary[f"K{index}"] = "配套子项，不另计分" if row.values["餐饮服务"] is not None and not result.cell_weights[row.display_name, "餐饮服务"] else "一级指标" if result.cell_weights[row.display_name, "餐饮服务"] else "--"
    formulas.set(summary, f"C{total_row}", f"=AVERAGE('周期计算'!C2:C{period_end})", result.aggregation.total_values["总分"])
    formulas.set(summary, f"I{total_row}", f"=SUM({tref('V')})", math.fsum(term.weight for term in result.types))
    formulas.set(summary, f"J{total_row}", f"=SUM({tref('X')})+SUM('周期计算'!N2:N{period_end})", math.fsum(term.rounding * term.weight for term in result.types) + result.aggregation.rounding_differences["总分"])
    summary[f"A{total_row + 1}"] = "一级指标权重"
    for col, name in enumerate(SUMMARY_COLUMNS[1:], 4):
        letter = get_column_letter(col)
        weight = math.fsum(term.weight for term in result.sections if term.column == name)
        formulas.set(summary, f"{letter}{total_row + 1}", f"=SUMIFS({ref('L')},{ref('F')},{letter}$2)", weight)
        formulas.set(summary, f"{letter}{total_row}", f'=IF({letter}{total_row + 1}=0,"",SUMIFS({ref("M")},{ref("F")},{letter}$2)/{letter}{total_row + 1})', result.column_scores[name] if weight else "")
    summary[f"A{total_row + 2}"] = "总分核对差额"
    formulas.set(summary, f"C{total_row + 2}", f"=C{total_row}-SUMPRODUCT(D{total_row}:H{total_row},D{total_row + 1}:H{total_row + 1})-J{total_row}", -result.checks["横向总分差额"])
    summary.merge_cells(start_row=total_row + 3, start_column=1, end_row=total_row + 3, end_column=11)
    summary.cell(total_row + 3, 1, "各周期等权。列底分数按一级指标贡献加权；餐饮为配套子项时权重为0。舍入差额保留单期原评分。")
    summary.row_dimensions[total_row + 3].height = 32
    summary.row_dimensions[1].height = 42
    summary.row_dimensions[2].height = 42
    for row in summary:
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if cell.column >= 9 or cell.row > total_row:
                cell.font = Font(name="Arial", size=11)
            if cell.data_type == "f":
                cell.number_format = "0.000000" if cell.column == 10 or (cell.column == 9 and cell.row > weights_header) else "0.00"
            if cell.row == total_row + 1 or (cell.row > weights_header and 3 <= cell.column <= 8) or (cell.column == 9 and 3 <= cell.row <= total_row):
                cell.number_format = "0.0000%"
    for col in ("I", "J", "K"):
        summary.column_dimensions[col].width = 22
        cell = summary[f"{col}2"]
        cell.font = copy(summary["C2"].font)
        cell.fill = copy(summary["C2"].fill)
        cell.border = copy(summary["C2"].border)
    summary[f"J{total_row}"].comment = Comment("类型舍入贡献之和 + 周期舍入贡献之和。详见周期计算及权重计算工作表。", "计算口径")
    style_calculation_sheet(cycle)
    style_calculation_sheet(detail)
    detail.auto_filter.ref = f"A1:O{end}"
    for sheet in workbook:
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = "1:2" if sheet is summary else "1:1"
        sheet.print_area = sheet.dimensions
    for sheet, coordinates in ((cycle, [f"C{row}" for row in range(2, trend_row + 1)]),):
        for coordinate in coordinates:
            sheet[coordinate].number_format = "0.00"
    path.parent.mkdir(parents=True, exist_ok=True)
    formulas.save(workbook, path)
    workbook.close()


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="从各独立周期原始数据重算，按周期等权生成满意度汇总及横纵核对表。")
    parser.add_argument("--year", required=True)
    parser.add_argument("--batch", required=True, help="输出批次，例如 1-9月 / Q1 / 1-12月")
    parser.add_argument("--periods", nargs="+", help="显式选择独立周期，例如 1-2月 3月 4月；不能包含累计批次。")
    parser.add_argument("--data-root", type=Path, default=Path("data"))
    parser.add_argument("--logs-root", type=Path, default=Path("logs/multi_period"))
    parser.add_argument("--config", type=Path, default=Path("pipeline.defaults.toml"))
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> None:
    args = parse_args(argv)
    try:
        year, batch = args.year.strip(), args.batch.strip()
        if not re.fullmatch(r"\d{4}", year) or not batch or Path(batch).name != batch or batch in {".", ".."}:
            raise ValueError("年份须为四位数字，批次须为单个目录名称。")
        labels = validate_period_labels(tuple(args.periods) if args.periods else period_labels_for_batch(batch))
        config = tomllib.loads(args.config.read_text(encoding="utf-8"))
        sheet_name = str(config.get("sheet_name", "问卷数据"))
        mode = normalize_calculation_mode(config.get("calculation_mode", "template"))
        periods: list[PeriodCalculation] = []
        for label in labels:
            period = calculate_period(
                args.data_root / "raw" / year / label, label,
                sheet_name=sheet_name, calculation_mode=mode,
                saved_summary_path=args.data_root / "satisfaction_summary" / year / label / f"{label}客户类型满意度汇总表.xlsx",
            )
            periods.append(period)
            print(f"{label}: {build_total_values(period.rows)['总分']:.2f}")
            for notice in period.notices:
                print(f"[提示] {label}: {notice}")
        result = build_multi_period_result(tuple(periods))
        output = args.data_root / "satisfaction_summary" / year / batch / f"{batch}客户类型满意度多周期汇总表.xlsx"
        write_multi_period_workbook(result, output, year=year, batch=batch)
        log = args.logs_root / year / batch / "calculation.json"
        log.parent.mkdir(parents=True, exist_ok=True)
        log.write_text(json.dumps({
            "year": year, "batch": batch, "periods": labels, "calculation_mode": mode,
            "source": "raw_recalculated", "total": result.aggregation.total_values["总分"],
            "period_totals": result.aggregation.period_totals,
            "period_rounding_contribution": result.aggregation.rounding_differences["总分"],
            "type_rounding_contribution": math.fsum(term.rounding * term.weight for term in result.types),
            "column_scores": result.column_scores, "checks": result.checks,
            "source_hashes": {period.label: period.source_hashes for period in periods},
            "notices": {period.label: period.notices for period in periods}, "output": str(output.resolve()),
        }, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(f"{len(periods)} 个周期等权总分：{result.aggregation.total_values['总分']:.2f}")
        print(f"汇总表：{output.resolve()}")
        print(f"计算日志：{log.resolve()}")
    except (ValueError, OSError) as error:
        raise SystemExit(f"多周期汇总失败：{error}") from error


if __name__ == "__main__":
    main()
