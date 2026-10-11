from __future__ import annotations

import hashlib
import json
import math
import re
import shutil
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from hangbo.summary.table import (
    SUMMARY_COLUMNS,
    ReportSnapshot,
    SummaryRowResult,
    build_summary_rows,
    build_total_values,
    load_summary_column_widths,
    style_summary_worksheet,
)
from hangbo.survey.customer_category_rules import DISPLAY_ORDERED_CUSTOMER_CATEGORY_RULES
from hangbo.survey.stats import (
    MetricResult,
    MetricDefinition,
    RoleDefinition,
    SectionDefinition,
    SectionResult,
    SurveyStatistics,
    build_result_dataframe,
    generate_customer_category_report_bundle,
    mean_ignore_empty,
    save_results,
)


MONTH_RANGE_RE = re.compile(r"(\d{1,2})(?:-(\d{1,2}))?月$")
QUARTER_RE = re.compile(r"Q([1-4])$", re.IGNORECASE)
VALID_JANUARY_FEBRUARY_MODES = {"auto", "grouped", "split"}


@dataclass(frozen=True)
class PeriodResult:
    """One independent period calculated through the existing survey engine."""

    label: str
    rows: tuple[SummaryRowResult, ...]
    reports: dict[str, SurveyStatistics] = field(default_factory=dict)
    raw_dir: Path | None = None
    source_hashes: dict[str, str] = field(default_factory=dict)
    notices: tuple[str, ...] = ()


@dataclass(frozen=True)
class CumulativeResult:
    periods: tuple[PeriodResult, ...]
    rows: tuple[SummaryRowResult, ...]
    total_values: dict[str, float | None]
    period_totals: dict[str, dict[str, float | None]]
    row_weights: dict[tuple[str, str], float]
    period_row_weights: dict[tuple[str, str, str], float]
    rounding_differences: dict[str, float]
    checks: dict[str, float]


def _parse_batch(batch: str) -> tuple[int, int]:
    normalized = str(batch).strip().replace("–", "-").replace("—", "-")
    quarter = QUARTER_RE.fullmatch(normalized)
    if quarter:
        start = 3 * int(quarter.group(1)) - 2
        return start, start + 2
    match = MONTH_RANGE_RE.fullmatch(normalized)
    if match is None:
        raise ValueError("批次须为月份、月份范围或 Q1-Q4。")
    start = int(match.group(1))
    end = int(match.group(2) or match.group(1))
    if not 1 <= start <= end <= 12:
        raise ValueError("月份范围须在 1-12 月内，且起始月份不大于结束月份。")
    return start, end


def is_cumulative_batch(batch: str) -> bool:
    normalized = str(batch).strip().replace("–", "-").replace("—", "-")
    if QUARTER_RE.fullmatch(normalized):
        return True
    match = MONTH_RANGE_RE.fullmatch(normalized)
    return bool(match and int(match.group(2) or match.group(1)) > int(match.group(1))) and normalized != "1-2月"


def _resolve_january_february_mode(
    raw_root: Path,
    mode: str,
) -> str:
    normalized = str(mode).strip().lower()
    if normalized not in VALID_JANUARY_FEBRUARY_MODES:
        raise ValueError(
            "january_february_mode 仅支持 auto、grouped 或 split。"
        )
    if normalized != "auto":
        return normalized
    has_grouped = (raw_root / "1-2月").is_dir()
    has_split = (raw_root / "1月").is_dir() and (raw_root / "2月").is_dir()
    if has_grouped and has_split:
        raise ValueError(
            "同时发现 1-2月、1月和2月目录，无法自动确定周期口径；"
            "请在 pipeline.defaults.toml 设置 january_february_mode = grouped 或 split。"
        )
    if has_split:
        return "split"
    if has_grouped:
        return "grouped"
    return "grouped"


def resolve_period_labels(
    batch: str,
    raw_root: Path,
    *,
    january_february_mode: str = "auto",
) -> tuple[str, ...]:
    start, end = _parse_batch(batch)
    if start <= 2:
        if start != 1 or end < 2:
            raise ValueError("包含 1 月或 2 月的累计批次必须从 1 月开始。")
        first_two = (
            ("1月", "2月")
            if _resolve_january_february_mode(raw_root, january_february_mode) == "split"
            else ("1-2月",)
        )
        return (*first_two, *(f"{month}月" for month in range(3, end + 1)))
    return tuple(f"{month}月" for month in range(start, end + 1))


def _source_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _build_snapshot(rule_name: str, source_path: Path, stats: SurveyStatistics) -> ReportSnapshot:
    return ReportSnapshot(
        role_name=rule_name,
        source_path=source_path,
        total=stats.satisfaction,
        sections={section.name: section.satisfaction for section in stats.sections},
        metrics={
            metric.name: metric.satisfaction
            for section in stats.sections
            for metric in section.metrics
        },
    )


def calculate_period(
    raw_dir: Path,
    label: str,
    *,
    sheet_name: str,
    calculation_mode: str,
) -> PeriodResult:
    """Calculate one period using the same report generator as single-period runs."""

    if not raw_dir.is_dir():
        raise ValueError(f"缺少所选周期原始数据目录：{raw_dir}")

    reports: dict[str, SurveyStatistics] = {}
    snapshots: list[ReportSnapshot] = []
    source_hashes: dict[str, str] = {}
    notices: list[str] = []

    with tempfile.TemporaryDirectory(prefix="hangbo-cumulative-period-") as temporary:
        working_dir = Path(temporary)
        for source_path in raw_dir.glob("*.xlsx"):
            shutil.copy2(source_path, working_dir / source_path.name)

        for rule in DISPLAY_ORDERED_CUSTOMER_CATEGORY_RULES:
            source_path = working_dir / rule.source_file_name
            original_path = raw_dir / rule.source_file_name
            if not source_path.is_file():
                notices.append(f"缺少来源文件：{rule.source_file_name}，相应客户类型无数据。")
                continue
            report = generate_customer_category_report_bundle(
                input_path=source_path,
                category_rule=rule,
                output_path=working_dir / f"{rule.name}.xlsx",
                sheet_name=sheet_name,
                calculation_mode=calculation_mode,
                dry_run=True,
                save_empty_report=False,
            )
            if report.stats.matched_row_count == 0:
                notices.append(f"{rule.customer_category}：无样本。")
                continue
            reports[rule.name] = report.stats
            snapshots.append(
                _build_snapshot(
                    rule.name,
                    raw_dir / f"{rule.name}.xlsx",
                    report.stats,
                )
            )
            source_hashes[rule.source_file_name] = _source_hash(original_path)
            if report.stats.satisfaction is None:
                notices.append(f"{rule.customer_category}：有样本，但无有效满意度回答。")

    rows = build_summary_rows(tuple(snapshots))
    if build_total_values(rows)["总分"] is None:
        raise ValueError(f"周期 {label} 没有有效满意度，已停止累计。")
    return PeriodResult(
        label=label,
        rows=rows,
        reports=reports,
        raw_dir=raw_dir,
        source_hashes=source_hashes,
        notices=tuple(notices),
    )


def _mean_defined(values: Iterable[float | None]) -> float | None:
    return mean_ignore_empty(list(values))


def build_cumulative_result(periods: tuple[PeriodResult, ...]) -> CumulativeResult:
    if not periods:
        raise ValueError("至少需要一个独立周期。")
    if len({period.label for period in periods}) != len(periods):
        raise ValueError("计算周期重复。")

    period_totals = {
        period.label: build_total_values(period.rows)
        for period in periods
    }
    if any(period_totals[period.label]["总分"] is None for period in periods):
        raise ValueError("所选周期存在无有效满意度的周期，不能静默减少分母。")

    definitions = periods[0].rows
    names = [(row.category_label, row.display_name) for row in definitions]
    if any(
        [(row.category_label, row.display_name) for row in period.rows] != names
        for period in periods
    ):
        raise ValueError("各期客户类型结构不一致。")

    total_values: dict[str, float | None] = {}
    row_weights: dict[tuple[str, str], float] = {}
    period_row_weights: dict[tuple[str, str, str], float] = {}
    contributions: dict[tuple[str, str], float] = {}
    checks: dict[str, float] = {}
    rounding_differences: dict[str, float] = {}

    for column in SUMMARY_COLUMNS:
        valid_periods = [
            period
            for period in periods
            if period_totals[period.label][column] is not None
        ]
        total_values[column] = _mean_defined(
            period_totals[period.label][column] for period in valid_periods
        )
        for period in valid_periods:
            valid_rows = [
                row
                for row in period.rows
                if row.display_name and row.values[column] is not None
            ]
            if not valid_rows:
                continue
            period_weight = 1 / len(valid_periods)
            for row in valid_rows:
                weight = period_weight / len(valid_rows)
                key = (row.display_name, column)
                period_row_weights[(period.label, row.display_name, column)] = weight
                row_weights[key] = row_weights.get(key, 0.0) + weight
                contributions[key] = contributions.get(key, 0.0) + weight * row.values[column]
        rounding_differences[column] = (total_values[column] or 0.0) - math.fsum(
            contributions[key]
            for key in contributions
            if key[1] == column
        )
        checks[f"{column}横向差额"] = (
            math.fsum(contributions[key] for key in contributions if key[1] == column)
            + rounding_differences[column]
            - (total_values[column] or 0.0)
        )

    rows = tuple(
        SummaryRowResult(
            category_label=row.category_label,
            display_name=row.display_name,
            values={
                column: (
                    contributions[(row.display_name, column)]
                    / row_weights[(row.display_name, column)]
                    if row.display_name and row_weights.get((row.display_name, column))
                    else None
                )
                for column in SUMMARY_COLUMNS
            },
            applicable_columns=row.applicable_columns,
        )
        for row in definitions
    )
    return CumulativeResult(
        periods=periods,
        rows=rows,
        total_values=total_values,
        period_totals=period_totals,
        row_weights=row_weights,
        period_row_weights=period_row_weights,
        rounding_differences=rounding_differences,
        checks=checks,
    )


def _weighted_average(
    values: Iterable[tuple[float | None, float]],
) -> float | None:
    defined = [(value, weight) for value, weight in values if value is not None]
    denominator = math.fsum(weight for _, weight in defined)
    if not defined or denominator == 0:
        return None
    return math.fsum(value * weight for value, weight in defined) / denominator


def _aggregate_stats_for_rule(
    result: CumulativeResult,
    rule_name: str,
    display_name: str,
) -> tuple[SurveyStatistics, RoleDefinition] | None:
    weighted_stats: list[tuple[SurveyStatistics, float]] = []
    for period in result.periods:
        stats = period.reports.get(rule_name)
        weight = result.period_row_weights.get((period.label, display_name, "总分"), 0.0)
        if stats is not None and weight and stats.satisfaction is not None:
            weighted_stats.append((stats, weight))
    if not weighted_stats:
        return None

    row = next(
        (
            row
            for row in result.rows
            if row.display_name == display_name
        ),
        None,
    )
    if row is None or row.values["总分"] is None:
        return None

    section_names: list[str] = []
    metric_names: dict[str, list[str]] = {}
    for stats, _ in weighted_stats:
        for section in stats.sections:
            if section.name not in section_names:
                section_names.append(section.name)
            metric_names.setdefault(section.name, [])
            for metric in section.metrics:
                if metric.name not in metric_names[section.name]:
                    metric_names[section.name].append(metric.name)

    sections: list[SectionResult] = []
    definitions: list[SectionDefinition] = []
    for section_name in section_names:
        section_values = [
            (section.satisfaction, weight)
            for stats, weight in weighted_stats
            for section in stats.sections
            if section.name == section_name
        ]
        importance_values = [
            (section.importance, weight)
            for stats, weight in weighted_stats
            for section in stats.sections
            if section.name == section_name
        ]
        metrics: list[MetricResult] = []
        metric_definitions: list = []
        for metric_name in metric_names[section_name]:
            metric_values = [
                (metric.satisfaction, weight)
                for stats, weight in weighted_stats
                for section in stats.sections
                if section.name == section_name
                for metric in section.metrics
                if metric.name == metric_name
            ]
            metric_importance = [
                (metric.importance, weight)
                for stats, weight in weighted_stats
                for section in stats.sections
                if section.name == section_name
                for metric in section.metrics
                if metric.name == metric_name
            ]
            metrics.append(
                MetricResult(
                    metric_name,
                    _weighted_average(metric_values),
                    _weighted_average(metric_importance),
                )
            )
            metric_definitions.append(MetricDefinition(metric_name, "A", "B"))
        sections.append(
            SectionResult(
                section_name,
                _weighted_average(section_values),
                _weighted_average(importance_values),
                tuple(metrics),
            )
        )
        definitions.append(SectionDefinition(section_name, tuple(metric_definitions)))

    overall_importance = _weighted_average(
        [(stats.importance, weight) for stats, weight in weighted_stats]
    )
    stats = SurveyStatistics(
        role_name=rule_name,
        satisfaction=row.values["总分"],
        importance=overall_importance,
        sections=tuple(sections),
        matched_row_count=sum(stats.matched_row_count for stats, _ in weighted_stats),
    )
    return stats, RoleDefinition(role_name=rule_name, role_column="A", sections=tuple(definitions))


def write_cumulative_detail_reports(result: CumulativeResult, output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    for rule in DISPLAY_ORDERED_CUSTOMER_CATEGORY_RULES:
        output_path = output_dir / f"{rule.name}.xlsx"
        aggregated = _aggregate_stats_for_rule(result, rule.name, rule.customer_category)
        if aggregated is None:
            if output_path.exists():
                output_path.unlink()
            continue
        stats, role_definition = aggregated
        save_results(
            build_result_dataframe(stats),
            output_path,
            role_definition,
            rule.name,
        )


def write_cumulative_summary_workbook(
    result: CumulativeResult,
    path: Path,
    *,
    year: str,
    batch: str,
) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "汇总表"
    style_summary_worksheet(
        summary,
        result.rows,
        column_widths=load_summary_column_widths(),
    )
    summary["A1"] = f"{year}年{batch}客户类型满意度累计汇总表"
    total_row = len(result.rows) + 3
    for index, column in enumerate(SUMMARY_COLUMNS, start=3):
        cell = summary.cell(total_row, index)
        cell.value = result.total_values[column]
        cell.number_format = "0.00"
    note_row = total_row + 2
    summary.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=8)
    summary.cell(
        note_row,
        1,
        "累计总分按独立周期等权平均；指标和客户类型按周期实际贡献加权，详见周期计算与权重核对。",
    )
    summary.cell(note_row, 1).alignment = Alignment(wrap_text=True, vertical="center")
    summary.cell(note_row, 1).font = Font(name="Arial", size=10, italic=True)

    cycle = workbook.create_sheet("周期计算")
    cycle.append(["周期", *SUMMARY_COLUMNS, "周期权重", "总分贡献"])
    period_weight = 1 / len(result.periods)
    for period in result.periods:
        totals = result.period_totals[period.label]
        cycle.append([
            period.label,
            *(totals[column] for column in SUMMARY_COLUMNS),
            period_weight,
            (totals["总分"] or 0) * period_weight,
        ])
    cycle.append(["各期均值", *(result.total_values[column] for column in SUMMARY_COLUMNS), 1.0, result.total_values["总分"]])

    weights = workbook.create_sheet("权重核对")
    weights.append(["周期", "样本类型", "指标", "单期值", "累计权重", "贡献"])
    for period in result.periods:
        for row in period.rows:
            if not row.display_name:
                continue
            for column in SUMMARY_COLUMNS:
                weight = result.period_row_weights.get((period.label, row.display_name, column), 0.0)
                if not weight:
                    continue
                value = row.values[column]
                weights.append([period.label, row.display_name, column, value, weight, value * weight])
    weights.append([])
    weights.append(["核对项", "差额"])
    for column, difference in result.rounding_differences.items():
        weights.append([f"{column}舍入差额", difference])
    for name, difference in result.checks.items():
        weights.append([name, difference])

    for sheet in (cycle, weights):
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            cell.fill = summary["A2"].fill.copy()
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.000000"
        for column in range(1, sheet.max_column + 1):
            sheet.column_dimensions[chr(64 + column) if column <= 26 else "A"].width = 18

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()


def run_cumulative_batch(
    *,
    year: str,
    batch: str,
    data_root: Path,
    detail_output_dir: Path,
    summary_output_path: Path,
    sheet_name: str,
    calculation_mode: str,
    january_february_mode: str,
    log_path: Path,
) -> CumulativeResult:
    raw_root = data_root / "raw" / year
    labels = resolve_period_labels(
        batch,
        raw_root,
        january_february_mode=january_february_mode,
    )
    periods = tuple(
        calculate_period(
            raw_root / label,
            label,
            sheet_name=sheet_name,
            calculation_mode=calculation_mode,
        )
        for label in labels
    )
    result = build_cumulative_result(periods)
    write_cumulative_detail_reports(result, detail_output_dir)
    write_cumulative_summary_workbook(
        result,
        summary_output_path,
        year=year,
        batch=batch,
    )
    log_path.parent.mkdir(parents=True, exist_ok=True)
    log_path.write_text(
        json.dumps(
            {
                "year": year,
                "batch": batch,
                "periods": labels,
                "period_mode": january_february_mode,
                "calculation_mode": calculation_mode,
                "source": "independent_periods_recalculated",
                "total": result.total_values["总分"],
                "period_totals": result.period_totals,
                "rounding_differences": result.rounding_differences,
                "checks": result.checks,
                "source_hashes": {
                    period.label: period.source_hashes for period in periods
                },
                "notices": {
                    period.label: period.notices for period in periods
                },
                "output": str(summary_output_path.resolve()),
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    return result
