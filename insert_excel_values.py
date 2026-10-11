"""Write TOML-configured values into an existing workbook without restyling it."""

from __future__ import annotations

import argparse
import math
import os
import re
import sys
import tempfile
import tomllib
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, MergedCell
from openpyxl.utils.cell import column_index_from_string
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet


Scalar = str | int | float | bool | date | datetime | time


@dataclass(frozen=True)
class CellWrite:
    sheet: str
    cell: str
    value: Scalar


@dataclass(frozen=True)
class ExcelJob:
    input_path: Path
    output_path: Path
    overwrite: bool
    keep_vba: bool
    writes: tuple[CellWrite, ...]


def check_table(
    value: Any, label: str, required: set[str], optional: set[str],
) -> dict[str, Any]:
    """Reject missing fields and spelling mistakes before opening a workbook."""
    if not isinstance(value, dict):
        raise ValueError(f"{label} 必须是 TOML 表")
    missing = required - value.keys()
    if missing:
        raise ValueError(f"{label} 缺少字段：{', '.join(sorted(missing))}")
    unknown = value.keys() - required - optional
    if unknown:
        raise ValueError(f"{label} 存在未知字段：{', '.join(sorted(unknown))}")
    return value


def require_string(value: Any, label: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{label} 必须是非空字符串")
    return value


def validate_value(value: Any, label: str) -> Scalar:
    if not isinstance(value, (str, int, float, bool, date, time)):
        raise ValueError(f"{label} 只支持文本、数值、布尔值和日期时间")
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError(f"{label} 不支持 nan 或 inf")
    if isinstance(value, str):
        if len(value) > 32767:
            raise ValueError(f"{label} 文本超过 Excel 的 32767 字符限制")
        if ILLEGAL_CHARACTERS_RE.search(value) or "\ufffe" in value or "\uffff" in value:
            raise ValueError(f"{label} 包含 Excel XML 不支持的字符")
    if isinstance(value, (datetime, time)) and value.tzinfo is not None:
        raise ValueError(f"{label} 不支持带时区的日期时间，请使用本地时间")
    return value


def load_job(config_path: Path) -> ExcelJob:
    """Parse a TOML job and normalize all cell addresses and file paths."""
    try:
        with config_path.open("rb") as stream:
            config = tomllib.load(stream)
    except tomllib.TOMLDecodeError as exc:
        raise ValueError(f"TOML 语法错误：{config_path}：{exc}") from exc
    except (OSError, UnicodeError) as exc:
        raise ValueError(f"无法读取配置文件：{config_path}：{exc}") from exc
    check_table(config, "配置", {"workbook", "writes"}, set())
    settings = check_table(
        config["workbook"], "workbook", {"input", "output"}, {"overwrite", "keep_vba"},
    )
    input_path = (config_path.parent / require_string(settings["input"], "workbook.input")).resolve()
    output_path = (config_path.parent / require_string(settings["output"], "workbook.output")).resolve()
    overwrite = settings.get("overwrite", False)
    keep_vba = settings.get("keep_vba", input_path.suffix.lower() == ".xlsm")
    for key, value in (("overwrite", overwrite), ("keep_vba", keep_vba)):
        if not isinstance(value, bool):
            raise ValueError(f"workbook.{key} 必须是布尔值 true 或 false")
    if input_path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("workbook.input 只支持 .xlsx 或 .xlsm 文件")
    if output_path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("workbook.output 只支持 .xlsx 或 .xlsm 文件")
    if input_path.suffix.lower() != output_path.suffix.lower():
        raise ValueError("workbook.input 和 workbook.output 的扩展名必须一致")
    if keep_vba != (input_path.suffix.lower() == ".xlsm"):
        raise ValueError(".xlsm 必须使用 keep_vba = true；.xlsx 必须使用 keep_vba = false")
    if not input_path.is_file():
        raise ValueError(f"workbook.input 文件不存在或不是文件：{input_path}")
    if not isinstance(config["writes"], list) or not config["writes"]:
        raise ValueError("writes 必须是非空的 [[writes]] 列表")
    writes = []
    targets = set()
    for index, raw_write in enumerate(config["writes"], start=1):
        label = f"writes[{index}]"
        write = check_table(raw_write, label, {"sheet", "cell", "value"}, set())
        sheet = require_string(write["sheet"], f"{label}.sheet")
        address = require_string(write["cell"], f"{label}.cell").strip()
        match = re.fullmatch(r"\$?([A-Za-z]{1,3})\$?([1-9][0-9]{0,6})", address)
        if not match:
            raise ValueError(f"{label}.cell 无效：{address}；请使用单个地址，例如 B2")
        column, row = match.groups()
        if column_index_from_string(column.upper()) > 16384 or int(row) > 1048576:
            raise ValueError(f"{label}.cell 超出 Excel 范围 A1:XFD1048576：{address}")
        cell = f"{column.upper()}{row}"
        target = (sheet, cell)
        if target in targets:
            raise ValueError(f"{label} 重复指定单元格：{sheet}!{cell}")
        targets.add(target)
        writes.append(CellWrite(sheet, cell, validate_value(write["value"], f"{label}.value")))
    return ExcelJob(input_path, output_path, overwrite, keep_vba, tuple(writes))


def save_workbook(workbook: Workbook, output_path: Path, *, overwrite: bool) -> None:
    """Finish a temporary workbook before publishing it to the destination."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        prefix=f".{output_path.stem}-", suffix=output_path.suffix,
        dir=output_path.parent, delete=False,
    ) as temporary:
        temporary_path = Path(temporary.name)
    try:
        workbook.save(temporary_path)
        if overwrite:
            os.replace(temporary_path, output_path)
        else:
            # A hard link publishes the completed file without replacing a file
            # that another process may have created after the initial check.
            try:
                os.link(temporary_path, output_path)
            except FileExistsError as exc:
                raise FileExistsError(
                    f"输出文件已存在：{output_path}；如需覆盖，请设置 overwrite = true"
                ) from exc
    finally:
        temporary_path.unlink(missing_ok=True)


def apply_config(config_path: Path) -> Path:
    """Apply the configured writes and return the resolved output path."""
    config_path = config_path.resolve()
    job = load_job(config_path)
    if job.output_path.exists() and not job.overwrite:
        raise FileExistsError(
            f"输出文件已存在：{job.output_path}；如需覆盖，请设置 overwrite = true"
        )
    try:
        workbook = load_workbook(
            job.input_path, data_only=False, keep_vba=job.keep_vba,
            keep_links=True, rich_text=True,
        )
    except Exception as exc:
        raise ValueError(f"读取 Excel 失败：{job.input_path}：{exc}") from exc
    try:
        targets = []
        for index, write in enumerate(job.writes, start=1):
            label = f"writes[{index}]"
            if write.sheet not in workbook.sheetnames:
                raise ValueError(f"{label}.sheet 工作表不存在：{write.sheet}")
            sheet = workbook[write.sheet]
            if not isinstance(sheet, Worksheet):
                raise ValueError(f"{label}.sheet 必须是单元格工作表：{write.sheet}")
            cell = sheet[write.cell]
            if isinstance(cell, MergedCell):
                raise ValueError(
                    f"{label}.cell {write.sheet}!{write.cell} 属于合并区域，请指定左上角单元格"
                )
            targets.append((cell, write.value))
        for cell, value in targets:
            number_format = cell.number_format
            cell.value = value
            # openpyxl assigns a date format automatically for date/time values.
            # Restore the template format even when it is General or numeric.
            cell.number_format = number_format
        save_workbook(workbook, job.output_path, overwrite=job.overwrite)
    finally:
        workbook.close()
        if workbook.vba_archive is not None:
            workbook.vba_archive.close()
    return job.output_path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="按 TOML 配置写入 Excel 单元格并保留原排版。")
    parser.add_argument("--config", type=Path, required=True, help="TOML 配置文件路径")
    args = parser.parse_args(argv)
    try:
        output_path = apply_config(args.config)
    except (OSError, ValueError, KeyError) as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 1
    print(f"完成：{output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
