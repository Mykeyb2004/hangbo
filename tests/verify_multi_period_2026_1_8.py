"""独立核对 2026 年 1-8 月多周期满意度。

该脚本只读取原始数据、单期结果、多周期结果和多周期计算日志，不调用
生产流水线，也不写入任何原始数据或已有结果表。

运行：
    UV_CACHE_DIR=/private/tmp/hangbo-satisfaction-uv-cache \
    uv run --no-sync python tests/verify_multi_period_2026_1_8.py
"""

from __future__ import annotations

import hashlib
import json
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
YEAR = "2026"
PERIODS = ("1-2月", *(f"{month}月" for month in range(3, 9)))
DATA_ROOT = ROOT / "data"
LOG_PATH = ROOT / "logs/multi_period/2026/1-8月/calculation.json"
OUTPUT_PATH = DATA_ROOT / "satisfaction_summary/2026/1-8月/1-8月客户类型满意度多周期汇总表.xlsx"


def display(value: Decimal) -> Decimal:
    return value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def read_total(path: Path) -> Decimal:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for row in workbook.active.iter_rows(values_only=True):
            if row and row[0] == "总分":
                return Decimal(str(row[2]))
    finally:
        workbook.close()
    raise AssertionError(f"未找到总分行：{path}")


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def verify() -> None:
    log = json.loads(LOG_PATH.read_text(encoding="utf-8"))
    assert tuple(log["periods"]) == PERIODS, log["periods"]
    assert "9月" not in log["periods"]
    assert len(log["periods"]) == 7

    saved_totals = []
    for period in PERIODS:
        path = DATA_ROOT / "satisfaction_summary" / YEAR / period / f"{period}客户类型满意度汇总表.xlsx"
        saved_totals.append(read_total(path))

    saved_mean = sum(saved_totals, Decimal(0)) / Decimal(len(saved_totals))
    recalculated_totals = [
        Decimal(str(log["period_totals"][period]["总分"])) for period in PERIODS
    ]
    recalculated_mean = sum(recalculated_totals, Decimal(0)) / Decimal(len(recalculated_totals))
    output_total = read_total(OUTPUT_PATH)

    # 历史 1-2 月表是 9.72，而当前原始数据重算是 9.73；显示值仍都汇总为 9.71。
    assert saved_totals == [
        Decimal("9.72"), Decimal("9.57"), Decimal("9.61"), Decimal("9.81"),
        Decimal("9.74"), Decimal("9.76"), Decimal("9.76"),
    ]
    assert recalculated_totals == [
        Decimal("9.73"), Decimal("9.57"), Decimal("9.61"), Decimal("9.81"),
        Decimal("9.74"), Decimal("9.76"), Decimal("9.76"),
    ]
    assert display(saved_mean) == Decimal("9.71")
    assert display(recalculated_mean) == Decimal("9.71")
    assert abs(output_total - recalculated_mean) < Decimal("1e-12")
    assert Decimal(str(log["total"])) == Decimal(str(float(recalculated_mean)))

    # 当前实现把周期总分和周期内各类汇总保留到表格关系中。
    checks = {name: abs(Decimal(str(value))) for name, value in log["checks"].items()}
    assert max(checks.values()) <= Decimal("1e-12"), checks

    # 生成表的公式和缓存值均存在；缓存总分、周期均值和总分核对差额一致。
    formula_book = load_workbook(OUTPUT_PATH, read_only=True, data_only=False)
    value_book = load_workbook(OUTPUT_PATH, read_only=True, data_only=True)
    try:
        assert formula_book.sheetnames == ["汇总表", "周期计算", "权重计算"]
        summary_formula = formula_book["汇总表"]
        summary_values = value_book["汇总表"]
        assert summary_formula["C22"].value.startswith("=AVERAGE(")
        assert abs(Decimal(str(summary_values["C22"].value)) - output_total) < Decimal("1e-12")
        assert abs(Decimal(str(summary_values["C24"].value))) <= Decimal("1e-12")

        cycle_formula = formula_book["周期计算"]
        cycle_values = value_book["周期计算"]
        assert [cycle_values[f"A{row}"].value for row in range(2, 9)] == list(PERIODS)
        assert cycle_formula["C9"].value.startswith("=AVERAGE(")
        assert abs(Decimal(str(cycle_values["C9"].value)) - recalculated_mean) < Decimal("1e-12")

        for formula_sheet, value_sheet in zip(formula_book.worksheets, value_book.worksheets):
            for row in formula_sheet.iter_rows():
                for formula_cell in row:
                    if formula_cell.data_type != "f":
                        continue
                    assert "#REF!" not in str(formula_cell.value)
                    cached = value_sheet[formula_cell.coordinate].value
                    assert cached is not None, (formula_sheet.title, formula_cell.coordinate)
    finally:
        formula_book.close()
        value_book.close()

    # 日志中的哈希必须与当前 raw 文件一致，证明计算过程没有改写原始文件。
    for period, files in log["source_hashes"].items():
        for filename, expected in files.items():
            source = DATA_ROOT / "raw" / YEAR / period / filename
            assert source.is_file(), source
            assert sha256(source) == expected, source

    cumulative_total = read_total(
        DATA_ROOT / "satisfaction_summary/2026/1-8月/1-8月客户类型满意度汇总表.xlsx"
    )
    print(f"周期数：{len(PERIODS)}（1-2月按一个周期）")
    print(f"已保存单期均值：{saved_mean} -> {display(saved_mean)}")
    print(f"原始数据重算均值：{recalculated_mean} -> {display(recalculated_mean)}")
    print(f"多周期表总分：{output_total} -> {display(output_total)}")
    print(f"同目录累计原始口径表总分：{cumulative_total}（不是本次多周期均值目标）")
    print(f"最大横纵核对差额：{max(checks.values())}")
    print("1-8月多周期满意度独立验证通过；原始文件哈希一致，未修改原始数据。")


if __name__ == "__main__":
    verify()
