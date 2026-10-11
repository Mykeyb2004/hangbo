from __future__ import annotations

import subprocess
import sys
import unittest
from copy import copy
from datetime import datetime, time
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.cell.text import InlineFont
from openpyxl.utils.datetime import to_excel

from insert_excel_values import apply_config


SCRIPT = Path(__file__).resolve().parents[1] / "insert_excel_values.py"


class InsertExcelValuesTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.root = Path(self.temp_dir.name)
        self.input_path = self.root / "source.xlsx"
        self.output_path = self.root / "result.xlsx"
        self.config_path = self.root / "insert.toml"

    def run_script(self, config_text: str) -> subprocess.CompletedProcess[str]:
        self.config_path.write_text(config_text, encoding="utf-8")
        return subprocess.run(
            [sys.executable, str(SCRIPT), "--config", str(self.config_path)],
            cwd=self.root.parent,
            capture_output=True,
            text=True,
            check=False,
        )

    def test_cli_writes_a_cell_and_preserves_layout_and_source(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "汇总"
        cell = sheet["B2"]
        cell.value = 1.25
        cell.font = Font(name="宋体", size=14, bold=True, color="FF990000")
        cell.fill = PatternFill("solid", fgColor="FFFFCCAA")
        cell.border = Border(bottom=Side(style="thin", color="FF000000"))
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.protection = Protection(locked=False)
        cell.number_format = "0.00"
        sheet.column_dimensions["B"].width = 24
        sheet.row_dimensions[2].height = 30
        sheet.merge_cells("D4:E4")
        sheet.freeze_panes = "B3"
        sheet.print_area = "A1:F12"
        sheet.page_setup.orientation = "landscape"
        sheet["F2"] = "=B2*2"
        workbook.create_sheet("明细")["A1"] = "保持"
        workbook.save(self.input_path)
        workbook.close()
        source_bytes = self.input_path.read_bytes()

        result = self.run_script('''[workbook]
input = "source.xlsx"
output = "result.xlsx"

[[writes]]
sheet = "汇总"
cell = "B2"
value = 12.5
''')

        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("完成", result.stdout)
        self.assertEqual(self.input_path.read_bytes(), source_bytes)
        saved = load_workbook(self.output_path)
        self.addCleanup(saved.close)
        source = load_workbook(self.input_path)
        self.addCleanup(source.close)
        self.assertEqual(saved["汇总"]["B2"].value, 12.5)
        for attribute in ("font", "fill", "border", "alignment", "protection", "number_format"):
            self.assertEqual(
                copy(getattr(saved["汇总"]["B2"], attribute)),
                copy(getattr(source["汇总"]["B2"], attribute)),
                attribute,
            )
        self.assertEqual(saved["汇总"].column_dimensions["B"].width, 24)
        self.assertEqual(saved["汇总"].row_dimensions[2].height, 30)
        self.assertEqual(str(saved["汇总"].merged_cells), "D4:E4")
        self.assertEqual(saved["汇总"].freeze_panes, "B3")
        self.assertEqual(saved["汇总"].print_area, "'汇总'!$A$1:$F$12")
        self.assertEqual(saved["汇总"].page_setup.orientation, "landscape")
        self.assertEqual(saved["汇总"]["F2"].value, "=B2*2")
        self.assertEqual(saved["明细"]["A1"].value, "保持")

    def make_workbook(self) -> None:
        workbook = Workbook()
        workbook.active.title = "汇总"
        workbook.active["B2"] = "原值"
        workbook.active.merge_cells("D4:E4")
        workbook.save(self.input_path)
        workbook.close()

    def config(self, *, output: str = "result.xlsx", settings: str = "", writes: str = "") -> str:
        return f'''[workbook]
input = "source.xlsx"
output = "{output}"
{settings}

[[writes]]
sheet = "汇总"
cell = "B2"
value = "新值"
{writes}
'''

    def test_existing_output_is_not_overwritten_by_default(self) -> None:
        self.make_workbook()
        self.output_path.write_bytes(b"existing destination")
        result = self.run_script(self.config())
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("overwrite", result.stderr)
        self.assertEqual(self.output_path.read_bytes(), b"existing destination")

    def test_explicit_overwrite_can_update_source_in_place(self) -> None:
        self.make_workbook()
        result = self.run_script(self.config(output="source.xlsx", settings="overwrite = true"))
        self.assertEqual(result.returncode, 0, result.stderr)
        saved = load_workbook(self.input_path)
        self.addCleanup(saved.close)
        self.assertEqual(saved["汇总"]["B2"].value, "新值")

    def test_source_cannot_be_overwritten_by_default(self) -> None:
        self.make_workbook()
        original = self.input_path.read_bytes()
        result = self.run_script(self.config(output="source.xlsx"))
        self.assertNotEqual(result.returncode, 0)
        self.assertEqual(self.input_path.read_bytes(), original)

    def test_output_parent_directories_are_created(self) -> None:
        self.make_workbook()
        result = self.run_script(self.config(output="nested/output/result.xlsx"))
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertTrue((self.root / "nested/output/result.xlsx").is_file())

    def test_invalid_configuration_is_rejected_without_output(self) -> None:
        self.make_workbook()
        valid = self.config()
        cases = [
            ("missing workbook", 'writes = []', "workbook"),
            ("missing input", valid.replace('input = "source.xlsx"\n', ""), "input"),
            ("missing output", valid.replace('output = "result.xlsx"\n', ""), "output"),
            ("invalid input", valid.replace('input = "source.xlsx"', "input = 123"), "input"),
            ("empty output", valid.replace('output = "result.xlsx"', 'output = ""'), "output"),
            ("invalid overwrite", self.config(settings='overwrite = "false"'), "overwrite"),
            ("invalid keep_vba", self.config(settings='keep_vba = "true"'), "keep_vba"),
            ("no writes", '[workbook]\ninput="source.xlsx"\noutput="result.xlsx"', "writes"),
            ("empty writes", 'writes=[]\n[workbook]\ninput="source.xlsx"\noutput="result.xlsx"', "writes"),
            ("invalid writes", 'writes=42\n[workbook]\ninput="source.xlsx"\noutput="result.xlsx"', "writes"),
            ("missing value", valid.replace('value = "新值"\n', ""), "value"),
            ("missing sheet", valid.replace('sheet = "汇总"\n', ""), "sheet"),
            ("missing cell", valid.replace('cell = "B2"\n', ""), "cell"),
            ("invalid sheet type", valid.replace('sheet = "汇总"', "sheet = 42"), "sheet"),
            ("invalid cell type", valid.replace('cell = "B2"', "cell = 42"), "cell"),
            ("invalid workbook table", "workbook=42\nwrites=[]", "workbook"),
            ("invalid write entry", 'writes=[42]\n[workbook]\ninput="source.xlsx"\noutput="result.xlsx"', "writes"),
            ("typo in setting", self.config(settings="overwirte = true"), "overwirte"),
            ("typo in write", self.config(writes="values = 42"), "values"),
        ]
        for name, config, expected in cases:
            with self.subTest(name=name):
                result = self.run_script(config)
                self.assertNotEqual(result.returncode, 0)
                self.assertIn(expected, result.stderr)
                self.assertIn("错误：", result.stderr)
                self.assertNotIn("Traceback", result.stderr)
                self.assertFalse(self.output_path.exists())

    def test_invalid_targets_leave_an_existing_destination_unchanged(self) -> None:
        self.make_workbook()
        original = self.input_path.read_bytes()
        self.output_path.write_bytes(b"existing destination")
        valid = self.config(settings="overwrite = true")
        invalid = [
            valid.replace('sheet = "汇总"', 'sheet = "不存在"'),
            valid.replace('cell = "B2"', 'cell = "E4"'),
            *[valid.replace('cell = "B2"', f'cell = "{cell}"')
              for cell in ("B0", "A1048577", "XFE1", "B2:C3", "2B", "A", "B-1", "Sheet1!B2")],
            valid + '\n[[writes]]\nsheet="汇总"\ncell="b2"\nvalue=123\n',
        ]
        for config in invalid:
            with self.subTest(config=config):
                result = self.run_script(config)
                self.assertNotEqual(result.returncode, 0)
                self.assertIn("writes", result.stderr)
                self.assertNotIn("Traceback", result.stderr)
                self.assertEqual(self.output_path.read_bytes(), b"existing destination")
                self.assertEqual(self.input_path.read_bytes(), original)

    def test_merged_anchor_and_absolute_lowercase_addresses_are_writable(self) -> None:
        self.make_workbook()
        result = self.run_script(self.config().replace('cell = "B2"', 'cell = "$d$4"'))
        self.assertEqual(result.returncode, 0, result.stderr)
        saved = load_workbook(self.output_path)
        self.addCleanup(saved.close)
        self.assertEqual(saved["汇总"]["D4"].value, "新值")
        self.assertEqual(str(saved["汇总"].merged_cells), "D4:E4")

    def test_invalid_values_are_rejected_without_output(self) -> None:
        self.make_workbook()
        invalid = ["[1, 2]", '{x = 1}', "nan", "inf", "-inf",
                   '"bad\\u0000text"', '"' + "x" * 32768 + '"',
                   "2026-10-09T08:30:00+08:00"]
        for value in invalid:
            with self.subTest(value=value[:60]):
                result = self.run_script(self.config().replace('value = "新值"', f"value = {value}"))
                self.assertNotEqual(result.returncode, 0)
                self.assertIn("writes", result.stderr)
                self.assertIn("value", result.stderr)
                self.assertNotIn("Traceback", result.stderr)
                self.assertFalse(self.output_path.exists())

    def test_multiple_sheets_and_scalar_types_keep_existing_number_formats(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "汇总"
        sheet["D2"].number_format = "yyyy/mm/dd"
        sheet["E2"].number_format = "hh:mm:ss"
        sheet["F2"].number_format = "yyyy/mm/dd hh:mm:ss"
        sheet["G2"].number_format = "0.00"
        workbook.create_sheet("明细")
        workbook.save(self.input_path)
        workbook.close()
        result = self.run_script('''[workbook]
input = "source.xlsx"
output = "result.xlsx"

[[writes]]
sheet = "汇总"
cell = "A2"
value = 0

[[writes]]
sheet = "汇总"
cell = "B2"
value = false

[[writes]]
sheet = "汇总"
cell = "C2"
value = ""

[[writes]]
sheet = "汇总"
cell = "D2"
value = 2026-10-09

[[writes]]
sheet = "汇总"
cell = "E2"
value = 08:30:00

[[writes]]
sheet = "汇总"
cell = "F2"
value = 2026-10-09T08:30:00

[[writes]]
sheet = "汇总"
cell = "G2"
value = 2026-10-09

[[writes]]
sheet = "明细"
cell = "A1"
value = "=汇总!A2+1"
''')
        self.assertEqual(result.returncode, 0, result.stderr)
        saved = load_workbook(self.output_path, data_only=False)
        self.addCleanup(saved.close)
        self.assertEqual(saved["汇总"]["A2"].value, 0)
        self.assertIs(saved["汇总"]["B2"].value, False)
        self.assertIsNone(saved["汇总"]["C2"].value)
        self.assertEqual(saved["汇总"]["D2"].value, datetime(2026, 10, 9))
        self.assertEqual(saved["汇总"]["E2"].value, time(8, 30))
        self.assertEqual(saved["汇总"]["F2"].value, datetime(2026, 10, 9, 8, 30))
        self.assertEqual(saved["汇总"]["G2"].number_format, "0.00")
        self.assertEqual(saved["汇总"]["G2"].value, to_excel(datetime(2026, 10, 9)))
        self.assertEqual(saved["明细"]["A1"].value, "=汇总!A2+1")

    def test_invalid_extensions_and_macro_settings_are_rejected(self) -> None:
        self.make_workbook()
        (self.root / "source.xls").write_bytes(self.input_path.read_bytes())
        (self.root / "source.xlsm").write_bytes(self.input_path.read_bytes())
        cases = [
            (self.config().replace("source.xlsx", "source.xls"), "xlsx"),
            (self.config(output="result.csv"), "output"),
            (self.config(output="result.xlsm"), "扩展名"),
            (self.config().replace("source.xlsx", "source.xlsm"), "扩展名"),
            (self.config(output="result.xlsm", settings="keep_vba = false")
             .replace("source.xlsx", "source.xlsm"), "keep_vba"),
            (self.config(settings="keep_vba = true"), "keep_vba"),
        ]
        for config, expected in cases:
            with self.subTest(config=config):
                result = self.run_script(config)
                self.assertNotEqual(result.returncode, 0)
                self.assertIn(expected, result.stderr)
                self.assertNotIn("Traceback", result.stderr)
                self.assertFalse(self.output_path.exists())
                self.assertFalse((self.root / "result.xlsm").exists())
                self.assertFalse((self.root / "result.csv").exists())

    def test_macro_payload_is_preserved_for_xlsm(self) -> None:
        self.make_workbook()
        with ZipFile(self.input_path) as archive:
            members = {name: archive.read(name) for name in archive.namelist()}
        members["[Content_Types].xml"] = members["[Content_Types].xml"].replace(
            b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
            b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
        ).replace(
            b"</Types>",
            b'<Override PartName="/xl/vbaProject.bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>',
        )
        members["xl/_rels/workbook.xml.rels"] = members["xl/_rels/workbook.xml.rels"].replace(
            b"</Relationships>",
            b'<Relationship Id="rIdVBA" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/></Relationships>',
        )
        # Opaque bytes test preservation only; no macro execution is involved.
        members["xl/vbaProject.bin"] = b"synthetic VBA preservation fixture"
        with ZipFile(self.root / "source.xlsm", "w") as archive:
            for name, contents in members.items():
                archive.writestr(name, contents)
        for settings in ("", "keep_vba = true"):
            with self.subTest(settings=settings):
                output = self.root / "result.xlsm"
                output.unlink(missing_ok=True)
                result = self.run_script(
                    self.config(output="result.xlsm", settings=settings)
                    .replace("source.xlsx", "source.xlsm")
                )
                self.assertEqual(result.returncode, 0, result.stderr)
                with ZipFile(output) as archive:
                    self.assertEqual(archive.read("xl/vbaProject.bin"), members["xl/vbaProject.bin"])
                saved = load_workbook(output, keep_vba=True)
                self.assertEqual(saved["汇总"]["B2"].value, "新值")
                saved.close()
                saved.vba_archive.close()

    def test_missing_and_malformed_files_report_errors_without_tracebacks(self) -> None:
        self.make_workbook()
        cases = [
            (self.config().replace("source.xlsx", "missing.xlsx"), "input"),
            ('[workbook\n', "TOML"),
        ]
        for config, expected in cases:
            with self.subTest(config=config):
                result = self.run_script(config)
                self.assertNotEqual(result.returncode, 0)
                self.assertIn(expected, result.stderr)
                self.assertIn("错误：", result.stderr)
                self.assertNotIn("Traceback", result.stderr)
                self.assertFalse(self.output_path.exists())
        self.input_path.write_bytes(b"invalid workbook")
        result = self.run_script(self.config())
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("Excel", result.stderr)
        self.assertNotIn("Traceback", result.stderr)
        self.assertFalse(self.output_path.exists())
        result = subprocess.run(
            [sys.executable, str(SCRIPT), "--config", str(self.root / "missing.toml")],
            capture_output=True, text=True, check=False,
        )
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("配置文件", result.stderr)
        self.assertNotIn("Traceback", result.stderr)

    def test_untouched_rich_text_is_preserved(self) -> None:
        workbook = Workbook()
        workbook.active.title = "汇总"
        workbook.active["A1"] = CellRichText(
            "普通文字", TextBlock(InlineFont(b=True, color="FFCC0000"), "加粗红色"),
        )
        workbook.save(self.input_path)
        workbook.close()
        result = self.run_script(self.config())
        self.assertEqual(result.returncode, 0, result.stderr)
        saved = load_workbook(self.output_path, rich_text=True)
        self.addCleanup(saved.close)
        value = saved["汇总"]["A1"].value
        self.assertIsInstance(value, CellRichText)
        self.assertEqual(str(value), "普通文字加粗红色")
        self.assertTrue(value[1].font.b)
        self.assertEqual(value[1].font.color.rgb, "FFCC0000")

    def test_save_failure_leaves_destination_intact_and_removes_temporary_file(self) -> None:
        self.make_workbook()
        original = self.input_path.read_bytes()
        self.output_path.write_bytes(b"existing destination")
        self.config_path.write_text(self.config(settings="overwrite = true"), encoding="utf-8")

        def fail_during_save(_workbook: Workbook, path: Path) -> None:
            Path(path).write_bytes(b"partial workbook")
            raise OSError("simulated disk write failure")

        # Inject a storage failure at the external workbook writer boundary.
        with patch.object(Workbook, "save", new=fail_during_save):
            with self.assertRaisesRegex(OSError, "disk write failure"):
                apply_config(self.config_path)
        self.assertEqual(self.output_path.read_bytes(), b"existing destination")
        self.assertEqual(self.input_path.read_bytes(), original)
        self.assertEqual(list(self.root.glob(".result-*")), [])

    def test_output_created_during_save_is_not_clobbered(self) -> None:
        self.make_workbook()
        self.config_path.write_text(self.config(), encoding="utf-8")
        original_save = Workbook.save

        def save_with_concurrent_destination(workbook: Workbook, path: Path) -> None:
            original_save(workbook, path)
            self.output_path.write_bytes(b"concurrently created file")

        with patch.object(Workbook, "save", new=save_with_concurrent_destination):
            with self.assertRaisesRegex(FileExistsError, "overwrite"):
                apply_config(self.config_path)
        self.assertEqual(self.output_path.read_bytes(), b"concurrently created file")
        self.assertEqual(list(self.root.glob(".result-*")), [])


if __name__ == "__main__":
    unittest.main()
