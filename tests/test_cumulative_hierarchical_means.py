"""Check cumulative mean relationships through the generated Excel workbooks."""

from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import load_workbook

from hangbo.pipeline.cumulative import (
    PeriodResult, build_cumulative_result, run_cumulative_batch, write_cumulative_detail_reports,
)
from hangbo.summary.table import SUMMARY_COLUMNS, SummaryRowResult, load_report_snapshots
from hangbo.survey.stats import MetricResult, SectionResult, SurveyStatistics


ROOT = Path(__file__).resolve().parents[1]


def rounded_mean(values):
    defined = [Decimal(str(value)) for value in values if value is not None]
    if not defined:
        return None
    return float((sum(defined) / len(defined)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


class CumulativeDetailRoundingTest(unittest.TestCase):
    def test_half_up_section_scores_and_empty_dimensions(self):
        # The summary reader rounds 2.675 to 2.68 (decimal half up), rather
        # than the 2.67 produced by binary float formatting.
        stats = SurveyStatistics(
            role_name="展览主承办", satisfaction=2.68, importance=None,
            sections=(
                SectionResult("会展服务", 2.675, None, (MetricResult("服务态度", 2.675, None),)),
                SectionResult("硬件设施", 2.69, None, (MetricResult("设施", 2.69, None),)),
                SectionResult("配套服务", None, None, (MetricResult("餐饮服务", None, None),)),
            ),
            matched_row_count=1,
        )
        values = dict.fromkeys(SUMMARY_COLUMNS)
        values["总分"] = stats.satisfaction
        row = SummaryRowResult("一、会展客户", "展览活动主（承）办", values, frozenset(SUMMARY_COLUMNS))
        result = build_cumulative_result((PeriodResult("7月", (row,), {"展览主承办": stats}),))
        with TemporaryDirectory() as directory:
            write_cumulative_detail_reports(result, Path(directory))
            report, = load_report_snapshots(Path(directory))
            self.assertEqual(report.sections, {"会展服务": 2.68, "硬件设施": 2.69, "配套服务": None})
            self.assertEqual(report.total, 2.69)


@unittest.skipUnless((ROOT / "data/raw/2026/9月/会展服务商.xlsx").is_file(), "local 2026 survey data required")
class CumulativeHierarchicalMeansTest(unittest.TestCase):
    def test_quarter_uses_published_dimension_means_and_customer_mean(self):
        with TemporaryDirectory() as directory:
            output = Path(directory)
            result = run_cumulative_batch(
                year="2026", batch="Q3", data_root=ROOT / "data",
                detail_output_dir=output / "details", summary_output_path=output / "summary.xlsx",
                sheet_name="问卷数据", calculation_mode="template",
                january_february_mode="auto", log_path=output / "cumulative.json",
            )
            reports = load_report_snapshots(output / "details")
            exhibition = next(report for report in reports if report.source_path.stem == "展览主承办")
            # Catering is nested under support service, not a fifth top-level dimension.
            self.assertEqual(sorted(exhibition.sections.values()), [9.55, 9.7, 9.77, 10.0])
            self.assertEqual(exhibition.total, 9.76)
            for report in reports:
                with self.subTest(customer=report.role_name):
                    self.assertEqual(report.total, rounded_mean(report.sections.values()))

            workbook = load_workbook(output / "summary.xlsx", data_only=True)
            try:
                summary = workbook["汇总表"]
                self.assertEqual(summary["C3"].value, 9.76)
                self.assertEqual(summary["C22"].value, 9.73)
                for col_index, indicator in enumerate(SUMMARY_COLUMNS, 3):
                    with self.subTest(indicator=indicator):
                        expected = rounded_mean(row.values[indicator] for row in result.rows if row.display_name)
                        self.assertEqual(summary.cell(22, col_index).value, expected)
                        self.assertEqual(result.total_values[indicator], expected)
                # Keep the distinct monthly reference visible without calling it the quarter result.
                self.assertEqual(workbook["周期计算"]["A5"].value, "客群等权汇总")
                self.assertEqual(workbook["周期计算"]["B5"].value, 9.73)
                self.assertEqual(workbook["周期计算"]["A6"].value, "各期均值（参考）")
                self.assertEqual(workbook["周期计算"]["B6"].value, 9.74)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
