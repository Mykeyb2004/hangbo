"""Read-only comparison of saved monthly/Q3 reports and Q3 merged raw rows.

Run: uv run --no-sync --no-cache python tests/verify_2026_monthly_q3_cross_tables.py
Only logs/verification/2026-monthly-q3-cross-tables.json is written.
This reports existing inconsistencies; success does not mean all tables agree.
"""

from collections import Counter
from decimal import Decimal, ROUND_HALF_UP
import hashlib
import json
from pathlib import Path
import re

from openpyxl import load_workbook

from hangbo.summary.table import SUMMARY_COLUMNS, build_summary_rows, load_report_snapshots


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
LABELS = tuple(f"{month}月" for month in range(3, 10)) + ("Q3",)


def normalized(value):
    if value is None or value == "--" or value == "":
        return None
    return Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def compare_reports(label):
    directory = DATA / "satisfaction_detail/2026" / label
    rebuilt = build_summary_rows(load_report_snapshots(directory))
    path = DATA / "satisfaction_summary/2026" / label / f"{label}客户类型满意度汇总表.xlsx"
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        saved = {
            row[1]: (index, row)
            for index, row in enumerate(workbook["汇总表"].values, 1)
            if len(row) >= 8 and row[1] in {r.display_name for r in rebuilt if r.display_name}
        }
        mismatches = []
        numeric_count = 0
        compared_count = 0
        for row in rebuilt:
            if not row.display_name:
                continue
            assert row.display_name in saved, (label, row.display_name, "missing summary row")
            index, saved_row = saved[row.display_name]
            for col_index, indicator in enumerate(SUMMARY_COLUMNS, 3):
                actual, detail = saved_row[col_index - 1], row.values[indicator]
                compared_count += 1
                numeric_count += normalized(actual) is not None or normalized(detail) is not None
                if normalized(actual) != normalized(detail):
                    mismatches.append({
                        "customer_type": row.display_name, "indicator": indicator,
                        "summary_cell": f"{chr(64 + col_index)}{index}",
                        "summary_value": actual, "from_detail_value": detail,
                    })
        return {"batch": label, "detail_workbooks": len(list(directory.glob("*.xlsx"))),
                "compared_cells": compared_count, "numeric_cells": numeric_count,
                "mismatches": mismatches}
    finally:
        workbook.close()


def read_raw(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["问卷数据"]
        iterator = iter(sheet.values)
        headers = ["" if v is None else str(v).strip() for v in next(iterator)]
        while headers and not headers[-1]:
            headers.pop()
        keys = tuple(re.sub(r"^Q\d+\s*[-_、.．]?", "", h, count=1).strip() for h in headers)
        assert len(keys) == len(set(keys)), (path.name, "duplicate semantic headers")
        rows = [tuple(row[:len(keys)]) for row in iterator
                if any(v is not None and (not isinstance(v, str) or v.strip()) for v in row[:len(keys)])]
        return keys, rows
    finally:
        workbook.close()


def compare_raw_merge():
    raw = DATA / "raw/2026"
    filenames = sorted({p.name for label in ("7月", "8月", "9月") for p in (raw / label).glob("*.xlsx")})
    results = []
    for filename in filenames:
        target_keys, target_rows = read_raw(raw / "Q3" / filename)
        expected = []
        counts = {}
        for label in ("7月", "8月", "9月"):
            path = raw / label / filename
            if not path.exists():
                counts[label] = 0
                continue
            keys, rows = read_raw(path)
            assert set(keys) <= set(target_keys), (filename, label, "missing columns")
            lookup = {key: i for i, key in enumerate(keys)}
            expected.extend(tuple(row[lookup[key]] if key in lookup else None for key in target_keys) for row in rows)
            counts[label] = len(rows)
        results.append({"workbook": filename, "monthly_rows": counts, "Q3_rows": len(target_rows),
                        "rows_and_values_match": Counter(expected) == Counter(target_rows),
                        "row_order_match": expected == target_rows})
    return results


if __name__ == "__main__":
    protected = set()
    for label in LABELS:
        for category in ("raw", "satisfaction_detail", "satisfaction_summary"):
            protected.update((DATA / category / "2026" / label).glob("*.xlsx"))
    before = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in protected}
    reports = [compare_reports(label) for label in LABELS]
    merged = compare_raw_merge()
    assert all(hashlib.sha256(p.read_bytes()).hexdigest() == h for p, h in before.items())
    output = ROOT / "logs/verification/2026-monthly-q3-cross-tables.json"
    output.parent.mkdir(parents=True, exist_ok=True)
    payload = {"reports": reports, "raw_merge": merged, "protected_excel_files_unchanged": len(protected)}
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({"reports": [{k: v for k, v in r.items() if k != "mismatches"} |
                                  {"mismatch_count": len(r["mismatches"])} for r in reports],
                      "raw_merge": merged, "protected_excel_files_unchanged": len(protected),
                      "audit_log": str(output)}, ensure_ascii=False, indent=2))
