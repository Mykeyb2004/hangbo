from __future__ import annotations

import unittest
import hashlib
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

from hangbo.multi_period_satisfaction import (
    aggregate_period_summary_rows,
    build_multi_period_result,
    calculate_period,
    period_labels_for_batch,
    validate_period_labels,
    write_multi_period_workbook,
    PeriodCalculation,
)
from hangbo.summary.table import SUMMARY_COLUMNS, SummaryRowResult
from hangbo.survey.stats import MetricResult, SectionResult, SurveyStatistics


def make_row(name: str, total: float | None, product: float | None) -> SummaryRowResult:
    values = {column: None for column in SUMMARY_COLUMNS}
    values["总分"] = total
    values["产品服务"] = product
    return SummaryRowResult(
        category_label="一、会展客户",
        display_name=name,
        values=values,
        applicable_columns=frozenset({"总分", "产品服务"}),
    )


def make_stats(name: str, total: float, sections: dict[str, float]) -> SurveyStatistics:
    return SurveyStatistics(
        name, total, None,
        tuple(SectionResult(label, score, None, (MetricResult(label, score, None),)) for label, score in sections.items()),
        1,
    )


def make_period(label: str, scores: dict[str, tuple[float, dict[str, float]]]) -> PeriodCalculation:
    rows = []
    components = {}
    for name in ("甲", "乙"):
        values = {column: None for column in SUMMARY_COLUMNS}
        if name in scores:
            total, sections = scores[name]
            values.update({"总分": total, "产品服务": sections.get("会展服务"), "硬件设施": sections.get("硬件设施")})
            components[name] = (make_stats(name, total, sections),)
        rows.append(SummaryRowResult("一、会展客户", name, values, frozenset(values)))
    return PeriodCalculation(label, tuple(rows), components, {}, (), None)


class MultiPeriodSatisfactionTest(unittest.TestCase):
    def test_full_year_can_have_11_or_12_registered_periods(self) -> None:
        self.assertEqual(len(period_labels_for_batch("1-12月")), 11)
        self.assertEqual(len(validate_period_labels(tuple(f"{month}月" for month in range(1, 13)))), 12)
        self.assertEqual(period_labels_for_batch("Q1"), ("1-2月", "3月"))
        self.assertEqual(period_labels_for_batch("Q4"), ("10月", "11月", "12月"))

    def test_reject_invalid_and_overlapping_periods(self) -> None:
        for labels in (("3月", "03月"), ("1-2月", "2月"), ("1-9月",), ()):
            with self.subTest(labels=labels), self.assertRaises(ValueError):
                validate_period_labels(labels)
        for batch in ("2-9月", "4-3月", "1-13月"):
            with self.subTest(batch=batch), self.assertRaises(ValueError):
                period_labels_for_batch(batch)

    def test_empty_period_is_not_silently_removed(self) -> None:
        with self.assertRaises(ValueError):
            aggregate_period_summary_rows({"p1": (make_row("甲", None, None),)})

    def test_hierarchy_reconciles_with_missing_sections_and_rounding(self) -> None:
        periods = (
            make_period("p1", {"甲": (9.02, {"会展服务": 9.01, "硬件设施": 9.02})}),
            make_period("p2", {"甲": (7.0, {"会展服务": 8.0, "硬件设施": 6.0}), "乙": (10.0, {"硬件设施": 10.0})}),
        )
        result = build_multi_period_result(periods)
        self.assertAlmostEqual(result.aggregation.total_values["总分"], 8.76)
        self.assertAlmostEqual(result.cell_weights["甲", "产品服务"], 0.375)
        self.assertAlmostEqual(result.cell_weights["乙", "产品服务"], 0)
        self.assertAlmostEqual(result.rows[0].values["产品服务"], (9.01 * 0.25 + 8 * 0.125) / 0.375)
        self.assertTrue(all(abs(value) < 1e-12 for value in result.checks.values()))
        self.assertAlmostEqual(sum(term.weight * term.rounding for term in result.types), 0.0025)

    def test_period_rounding_is_separate_from_category_scores(self) -> None:
        result = build_multi_period_result((make_period("p1", {
            "甲": (9.01, {"会展服务": 9.01}), "乙": (9.02, {"会展服务": 9.02}),
        }),))
        self.assertEqual(result.aggregation.total_values["总分"], 9.02)
        self.assertAlmostEqual(result.aggregation.rounding_differences["总分"], 0.005)
        self.assertEqual(result.rows[0].values["总分"], 9.01)
        self.assertTrue(all(abs(value) < 1e-12 for value in result.checks.values()))

    def test_hotel_components_keep_each_components_section_weights(self) -> None:
        row = make_row("甲", 8.0, None)
        components = (
            make_stats("宴会组件", 10, {"餐饮服务": 10, "硬件设施": 10}),
            make_stats("自助餐组件", 6, {"餐饮服务": 6}),
        )
        period = PeriodCalculation("p1", (row,), {"甲": components}, {}, (), None)
        result = build_multi_period_result((period,))
        self.assertAlmostEqual(result.cell_weights["甲", "餐饮服务"], 0.75)
        self.assertAlmostEqual(result.rows[0].values["餐饮服务"], 22 / 3)
        self.assertAlmostEqual(result.cell_weights["甲", "硬件设施"], 0.25)
        self.assertTrue(all(abs(value) < 1e-12 for value in result.checks.values()))

    def test_dining_subitem_has_zero_scoring_weight(self) -> None:
        values = {column: None for column in SUMMARY_COLUMNS}
        values.update({"总分": 8.0, "配套服务": 8.0, "餐饮服务": 6.0})
        row = SummaryRowResult("一、会展客户", "甲", values, frozenset(values))
        component = make_stats("甲", 8, {"配套服务": 8})
        period = PeriodCalculation("p1", (row,), {"甲": (component,)}, {}, (), None)
        result = build_multi_period_result((period,))
        self.assertEqual(result.rows[0].values["餐饮服务"], 6)
        self.assertEqual(result.cell_weights["甲", "餐饮服务"], 0)
        self.assertIsNone(result.column_scores["餐饮服务"])
        self.assertAlmostEqual(result.aggregation.total_values["总分"], 8)

    def test_excel_contains_formulas_cached_scores_and_checks(self) -> None:
        result = build_multi_period_result((
            make_period("p1", {"甲": (9.01, {"会展服务": 9.01}), "乙": (9.02, {"会展服务": 9.02})}),
            make_period("p2", {"甲": (6, {"会展服务": 6})}),
        ))
        with TemporaryDirectory() as directory:
            output = Path(directory) / "result.xlsx"
            write_multi_period_workbook(result, output, year="2026", batch="测试")
            workbook = load_workbook(output, data_only=False)
            try:
                self.assertEqual(workbook.sheetnames, ["汇总表", "周期计算", "权重计算"])
                self.assertIn("AVERAGE", workbook["汇总表"]["C5"].value)
                self.assertIn("SUMPRODUCT", workbook["汇总表"]["C7"].value)
            finally:
                workbook.close()
            cached = load_workbook(output, data_only=True)
            try:
                self.assertAlmostEqual(cached["汇总表"]["C5"].value, 7.51)
                self.assertAlmostEqual(cached["汇总表"]["C7"].value, 0)
                self.assertAlmostEqual(cached["汇总表"]["I12"].value, 0)
                self.assertIsNone(cached["汇总表"]["E5"].value)
                self.assertFalse(any(cell.data_type == "e" for sheet in cached for row in sheet for cell in row))
            finally:
                cached.close()

    def test_raw_workbook_preprocessing_only_changes_temporary_copy(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "餐饮.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "问卷数据"
            sheet.append([f"字段{index}" for index in range(1, 101)])
            sheet.append([1, "有效", "一期", "餐饮", "商务简餐", *([10] * 95)])
            workbook.save(source)
            workbook.close()
            before = hashlib.sha256(source.read_bytes()).hexdigest()
            period = calculate_period(root, "3月")
            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), before)
            self.assertEqual(period.source_hashes["餐饮.xlsx"], before)
            self.assertTrue(any("临时副本" in message for message in period.notices))
            self.assertEqual(next(row.values["总分"] for row in period.rows if row.display_name == "商务简餐"), 10)
        with self.assertRaises(ValueError):
            calculate_period(root, "3月")
    def test_period_labels_keep_1_2_month_as_one_period(self) -> None:
        self.assertEqual(
            period_labels_for_batch("1-9月"),
            ("1-2月", "3月", "4月", "5月", "6月", "7月", "8月", "9月"),
        )
        self.assertEqual(period_labels_for_batch("1-2月"), ("1-2月",))

    def test_aggregate_rows_reconciles_to_equal_period_total(self) -> None:
        period_rows = {
            "p1": (make_row("甲", 10.0, 10.0), make_row("乙", 8.0, 8.0)),
            "p2": (make_row("甲", 6.0, 6.0), make_row("乙", None, None)),
        }

        result = aggregate_period_summary_rows(period_rows)

        self.assertAlmostEqual(result.total_values["总分"], 7.5)
        self.assertAlmostEqual(result.total_values["产品服务"], 7.5)
        self.assertAlmostEqual(result.rows[0].values["总分"], 22 / 3)
        self.assertAlmostEqual(result.rows[1].values["总分"], 8.0)
        self.assertAlmostEqual(
            sum(
                result.row_weights[(row.display_name, "总分")]
                * (row.values["总分"] or 0)
                for row in result.rows
                if row.display_name
            ),
            result.total_values["总分"],
        )


if __name__ == "__main__":
    unittest.main()
