"""Compare the proposed first-level mean rule with July-September 2026.

Run: uv run python tests/verify_monthly_summary_from_details.py
Only logs/verification/2026-monthly-summary-from-details.json is written.
"""

from dataclasses import replace
from decimal import Decimal, ROUND_HALF_UP
import hashlib
import json
from pathlib import Path
import shutil
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from hangbo.summary.table import (
    SUMMARY_COLUMNS,
    build_summary_rows,
    build_total_values,
    load_report_snapshots,
)
from hangbo.survey.customer_category_rules import CUSTOMER_CATEGORY_RULE_BY_NAME
from hangbo.survey.stats import generate_customer_category_report_bundle


ROOT = Path(__file__).resolve().parents[1]
LABELS = ("7月", "8月", "9月")
protected = set()
for label in LABELS:
    protected.update((ROOT / "data/satisfaction_detail/2026" / label).glob("*.xlsx"))
    protected.add(ROOT / f"data/satisfaction_summary/2026/{label}/{label}客户类型满意度汇总表.xlsx")
protected.update(ROOT / p for p in (
    "data/raw/2026/8月/餐饮.xlsx", "src/hangbo/survey/stats.py", "src/hangbo/summary/table.py",
))
before = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in protected}
results = []
for label in LABELS:
    snapshots = load_report_snapshots(ROOT / "data/satisfaction_detail/2026" / label)
    corrected = []
    for snapshot in snapshots:
        defined = [Decimal(str(v)) for v in snapshot.sections.values() if v is not None]
        total = float((sum(defined) / len(defined)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)) if defined else None
        corrected.append(replace(snapshot, total=total))
    rows = build_summary_rows(tuple(corrected))
    totals = build_total_values(rows)
    path = ROOT / f"data/satisfaction_summary/2026/{label}/{label}客户类型满意度汇总表.xlsx"
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["汇总表"]
        lookup = {sheet.cell(i, 2).value: i for i in range(3, 22) if sheet.cell(i, 2).value}
        differences = []
        checks = []
        for row in rows:
            if not row.display_name:
                continue
            for j, column in enumerate(SUMMARY_COLUMNS, 3):
                cell = sheet.cell(lookup[row.display_name], j)
                old = None if cell.value in (None, "--", "") else cell.value
                checks.append({"customer_type": row.display_name, "indicator": column,
                               "cell": cell.coordinate, "current": old, "proposed": row.values[column]})
        for j, column in enumerate(SUMMARY_COLUMNS, 3):
            cell = sheet.cell(22, j)
            checks.append({"customer_type": "表底整体", "indicator": column,
                           "cell": cell.coordinate, "current": cell.value, "proposed": totals[column]})
        differences = [c for c in checks if c["current"] != c["proposed"]]
        results.append({"month": label, "detail_reports": len(snapshots), "compared_cells": len(checks),
                        "differences": differences, "bottom_totals": totals, "comparisons": checks})
    finally:
        workbook.close()

assert results[0]["differences"] == []
assert results[1]["differences"] == [{"customer_type": "酒店餐饮客户", "indicator": "总分",
                                      "cell": "C21", "current": 9.97, "proposed": 9.98}]
assert results[2]["differences"] == []

# Reproduce the existing aggregate on a temporary copy of the raw workbook.
rule = CUSTOMER_CATEGORY_RULE_BY_NAME["酒店餐饮客户"]
components = []
with TemporaryDirectory(prefix="hangbo-monthly-audit-") as temporary:
    temp = Path(temporary)
    raw_copy = temp / "餐饮.xlsx"
    shutil.copy2(ROOT / "data/raw/2026/8月/餐饮.xlsx", raw_copy)
    for name in rule.aggregate_rule_names:
        report = generate_customer_category_report_bundle(
            raw_copy, CUSTOMER_CATEGORY_RULE_BY_NAME[name], temp / f"{name}.xlsx",
            sheet_name="问卷数据", calculation_mode="template", dry_run=True,
        )
        components.append({"name": name, "sample_n": report.stats.matched_row_count,
                           "total": report.stats.satisfaction,
                           "sections": {s.name: s.satisfaction for s in report.stats.sections}})
    report = generate_customer_category_report_bundle(
        raw_copy, rule, temp / "aggregate.xlsx", sheet_name="问卷数据",
        calculation_mode="template", dry_run=True,
    )
    assert report.stats.satisfaction == 9.97
    assert {s.name: s.satisfaction for s in report.stats.sections} == {
        "餐饮服务": 9.95, "硬件设施": 10.0, "智慧场馆": None,
    }
for path, expected in before.items():
    assert hashlib.sha256(path.read_bytes()).hexdigest() == expected, path

payload = {"months": results, "august_hotel_dining_components": components,
           "protected_files_unchanged": len(before),
           "total_compared_cells": sum(r["compared_cells"] for r in results),
           "total_changed_cells": sum(len(r["differences"]) for r in results)}
output = ROOT / "logs/verification/2026-monthly-summary-from-details.json"
output.parent.mkdir(parents=True, exist_ok=True)
output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps({"months": [{k: v for k, v in r.items() if k != "comparisons"} for r in results],
                  "total_compared_cells": payload["total_compared_cells"],
                  "total_changed_cells": payload["total_changed_cells"],
                  "protected_files_unchanged": len(before)}, ensure_ascii=False, indent=2))
