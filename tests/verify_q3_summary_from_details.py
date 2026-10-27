"""Simulate summary/detail alignment without modifying source code or Excel.

uv run python tests/verify_q3_summary_from_details.py
"""

from dataclasses import replace
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

from hangbo.summary.table import (
    build_summary_rows,
    build_total_values,
    load_report_snapshots,
)
from hangbo.survey.stats import mean_ignore_empty


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
protected = list((DATA / "satisfaction_detail/2026/Q3").glob("*.xlsx"))
protected += list((DATA / "satisfaction_detail/2026/9月").glob("*.xlsx"))
summary_path = DATA / "satisfaction_summary/2026/Q3/Q3客户类型满意度汇总表.xlsx"
protected.append(summary_path)
before = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in protected}

# Follow the monthly convention: round each section to two decimal places,
# average only defined first-level sections, then round the overall score.
# load_report_snapshots already applies the existing two-decimal convention.
september = load_report_snapshots(DATA / "satisfaction_detail/2026/9月")
for snapshot in september:
    assert snapshot.total == mean_ignore_empty(list(snapshot.sections.values())), snapshot.role_name

quarter = load_report_snapshots(DATA / "satisfaction_detail/2026/Q3")
# For these six exhibition/conference customer types, dining is a child
# metric of supporting services, never an additional first-level section.
exhibition_roles = {"展览主承办", "参展商", "专业观众", "会展服务商", "会议主承办", "参会人员"}
assert exhibition_roles.issubset({s.role_name for s in quarter})
for snapshot in quarter:
    if snapshot.role_name in exhibition_roles:
        assert "餐饮服务" not in snapshot.sections, snapshot.role_name
        assert "餐饮服务" in snapshot.metrics, snapshot.role_name
corrected = tuple(replace(s, total=mean_ignore_empty(list(s.sections.values()))) for s in quarter)
old_rows = build_summary_rows(quarter)
new_rows = build_summary_rows(corrected)
for old, new in zip(old_rows, new_rows, strict=True):
    assert old.display_name == new.display_name
    for column in old.values:
        if column != "总分":
            assert old.values[column] == new.values[column], (old.display_name, column)
for row in new_rows[:6]:
    assert row.values["总分"] == mean_ignore_empty([
        row.values[c] for c in ("产品服务", "硬件设施", "配套服务", "智慧场馆/服务")
    ]), row.display_name

supplier = next(r for r in new_rows if r.display_name == "会展服务商")
assert supplier.values["总分"] == 9.85
assert supplier.values["智慧场馆/服务"] == 9.84
assert supplier.values["总分"] == mean_ignore_empty([
    supplier.values[c] for c in ("产品服务", "硬件设施", "配套服务", "智慧场馆/服务")
])
# Dining is already inside supporting services for this customer type.
assert supplier.values["餐饮服务"] == 9.42
assert mean_ignore_empty([supplier.values[c] for c in (
    "产品服务", "硬件设施", "配套服务", "智慧场馆/服务", "餐饮服务"
)]) != supplier.values["总分"]

workbook = load_workbook(summary_path, read_only=True, data_only=True)
try:
    sheet = workbook["汇总表"]
    old_saved = {sheet.cell(i, 2).value: sheet.cell(i, 3).value for i in range(3, 22) if sheet.cell(i, 2).value}
finally:
    workbook.close()

payload = {
    "basis": "Keep quarterly detail first-level values at monthly two-decimal precision; recompute totals from defined first-level sections; copy to summary",
    "september_reports_reconciled": len(september),
    "quarterly_reports": len(quarter),
    "exhibition_customer_types_without_dining_as_first_level": 6,
    "quarterly_report_totals_changed": sum(s.total != n.total for s, n in zip(quarter, corrected, strict=True)),
    "rows": [{"customer_type": r.display_name, "current_summary_total": old_saved[r.display_name], **r.values}
             for r in new_rows if r.display_name],
    "bottom_totals": build_total_values(new_rows),
}
for path, expected in before.items():
    assert hashlib.sha256(path.read_bytes()).hexdigest() == expected, path
payload["protected_excel_files_unchanged"] = len(before)
output = ROOT / "logs/verification/2026-Q3-summary-from-details.json"
output.parent.mkdir(parents=True, exist_ok=True)
output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(payload, ensure_ascii=False, indent=2))
