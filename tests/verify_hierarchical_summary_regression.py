"""Verify monthly preservation and cumulative hierarchy against the saved baseline.

Run with: uv run python tests/verify_hierarchical_summary_regression.py
Only temporary workbooks are generated; production Excel files are protected by hashes.
"""

from contextlib import redirect_stdout
from decimal import Decimal, ROUND_HALF_UP
import io
import json
from pathlib import Path
import shutil
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from hangbo.pipeline.cumulative import run_cumulative_batch
from hangbo.summary.table import SUMMARY_COLUMNS, build_summary_rows, generate_summary_report, load_report_snapshots
from hangbo.survey.stats import run_directory_batch
from verify_summary_algorithm_regression import (
    BASELINE, CONFIG, LABELS, ROOT, detail_snapshots, digest, saved_month, snapshot,
)


def mean(values):
    defined = [Decimal(str(value)) for value in values if value is not None]
    if not defined:
        return None
    return float((sum(defined) / len(defined)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def main():
    baseline = json.loads(BASELINE.read_text())
    protected = {
        path: digest(path)
        for label in (*LABELS, "Q3", "1-9月")
        for category in ("raw", "satisfaction_detail", "satisfaction_summary")
        for path in (ROOT / "data" / category / "2026" / label).glob("*.xlsx")
    }
    monthly_count = 0
    cumulative_checks = {}
    with TemporaryDirectory(prefix="hangbo-hierarchy-regression-") as directory:
        work = Path(directory)
        for label in LABELS:
            raw = work / "raw" / label
            shutil.copytree(ROOT / "data/raw/2026" / label, raw)
            detail = work / "detail" / label
            with redirect_stdout(io.StringIO()):
                run_directory_batch(input_dir=raw, output_dir=detail,
                                    sheet_name=CONFIG["sheet_name"], calculation_mode=CONFIG["calculation_mode"])
            summary = generate_summary_report(input_dir=detail, output_dir=work / "summary" / label,
                                             output_name=f"{label}客户类型满意度汇总表.xlsx")
            current = {"details": detail_snapshots(detail), "summary": snapshot(summary)}
            assert current == baseline["monthly"][label], (label, "monthly values changed from pre-change baseline")
            assert current == saved_month(label), (label, "monthly values differ from published files")
            monthly_count += len(current["details"])
            print(f"{label}: {len(current['details'])} detail workbooks and summary unchanged", flush=True)

        for label in ("Q3", "1-9月"):
            detail = work / "cumulative" / label / "detail"
            summary = work / "cumulative" / label / "summary.xlsx"
            with redirect_stdout(io.StringIO()):
                result = run_cumulative_batch(
                    year="2026", batch=label, data_root=ROOT / "data", detail_output_dir=detail,
                    summary_output_path=summary, sheet_name=CONFIG["sheet_name"],
                    calculation_mode=CONFIG["calculation_mode"], january_february_mode=CONFIG["january_february_mode"],
                    log_path=work / "cumulative" / label / "cumulative.json",
                )
            # Customer satisfaction (B2) is the sole permitted detail-cell change.
            current_details = detail_snapshots(detail)
            previous_details = baseline["cumulative"][label]["details"]
            assert current_details.keys() == previous_details.keys(), (label, "detail file set changed")
            for filename, sheets in current_details.items():
                assert sheets.keys() == previous_details[filename].keys(), (label, filename, "sheet set changed")
                for sheet, cells in sheets.items():
                    assert {cell: value for cell, value in cells.items() if cell != "B2"} == {
                        cell: value for cell, value in previous_details[filename][sheet].items() if cell != "B2"
                    }, (label, filename, "non-total detail values changed")

            reports = load_report_snapshots(detail)
            for report in reports:
                assert report.total == mean(report.sections.values()), (label, report.role_name, "dimension mean mismatch")
            rows = build_summary_rows(reports)
            workbook = load_workbook(summary, data_only=True)
            comparisons = 0
            try:
                sheet = workbook["汇总表"]
                for row_index, row in enumerate(rows, 3):
                    for col_index, indicator in enumerate(SUMMARY_COLUMNS, 3):
                        value = sheet.cell(row_index, col_index).value
                        actual = value if isinstance(value, (int, float)) else None
                        assert actual == row.values[indicator], (label, row.display_name, indicator, "cross-table mismatch")
                        comparisons += 1
                for col_index, indicator in enumerate(SUMMARY_COLUMNS, 3):
                    expected = mean(row.values[indicator] for row in rows if row.display_name)
                    assert sheet.cell(len(rows) + 3, col_index).value == expected, (label, indicator, "customer mean mismatch")
            finally:
                workbook.close()
            if label == "Q3":
                assert result.total_values["总分"] == 9.73
            cumulative_checks[label] = {
                "customer_totals_equal_dimension_means": len(reports),
                "cross_table_cells_checked": comparisons,
                "bottom_customer_means_checked": len(SUMMARY_COLUMNS),
                "non_total_detail_cells_unchanged": True,
                "overall_satisfaction": result.total_values["总分"],
                "monthly_overall_mean_reference": mean(t["总分"] for t in result.period_totals.values()),
            }
            print(f"{label}: {len(reports)} customer means, {comparisons} cross-table cells; total={result.total_values['总分']}", flush=True)

    assert all(digest(path) == expected for path, expected in protected.items()), "Production Excel files changed"
    report = {
        "monthly_labels": list(LABELS), "monthly_detail_workbooks_unchanged": monthly_count,
        "monthly_summary_workbooks_unchanged": len(LABELS),
        "monthly_values_match_pre_change_baseline": True,
        "cumulative": cumulative_checks, "protected_excel_files_unchanged": len(protected),
    }
    output = ROOT / "logs/verification/hierarchical-summary-regression.json"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
