"""Read-only audit of the 2026 Q3 smart-venue discrepancy.

Run from the repository root:
    uv run python tests/verify_q3_smart_venue_consistency.py

Only the audit JSON under logs/verification is written. Known inconsistent
outputs are reproduced and reported, not accepted as a consistency test.
"""

from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import hashlib
import json
import math
from pathlib import Path
import tomllib

from openpyxl import load_workbook

from hangbo.pipeline.cumulative import (
    _aggregate_stats_for_rule,
    build_cumulative_result,
    calculate_period,
)
from hangbo.summary.table import (
    SUMMARY_COLUMNS,
    build_summary_rows,
    load_report_snapshots,
)


ROOT = Path(__file__).resolve().parents[1]
LABELS = ("7月", "8月", "9月")
RAW = ROOT / "data/raw/2026"
DETAIL = ROOT / "data/satisfaction_detail/2026/Q3"
SUMMARY = ROOT / "data/satisfaction_summary/2026/Q3/Q3客户类型满意度汇总表.xlsx"
CONFIG = tomllib.loads((ROOT / "pipeline.defaults.toml").read_text())
RUN_LOG = json.loads((ROOT / "logs/pipeline/2026/Q3/cumulative.json").read_text())


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def close(actual, expected):
    assert math.isclose(actual, expected, rel_tol=0, abs_tol=1e-12), (actual, expected)


def rounded(value):
    return float(Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def raw_smart_values(label):
    """Independent header-based calculation, without the production stats engine."""
    workbook = load_workbook(RAW / label / "会展服务商.xlsx", read_only=True, data_only=True)
    try:
        rows = iter(workbook[CONFIG["sheet_name"]].values)
        headers = next(rows)
        selected = [r for r in rows if str(r[3]).strip() == "会展服务商"]
        metrics = []
        for question in ("Q19-", "Q21-", "Q17-", "Q23-"):
            indices = [i for i, h in enumerate(headers) if str(h).startswith(question) and "满意度" in str(h)]
            assert len(indices) == 1, (label, question, indices)
            values = []
            for row in selected:
                try:
                    value = Decimal(str(row[indices[0]]))
                except InvalidOperation:
                    continue
                if value.is_finite() and 0 < value < 11:
                    values.append(value)
            mean = rounded(sum(values) / len(values)) if values else None
            metrics.append({"question": question, "valid_n": len(values), "sum": float(sum(values)), "mean": mean})
        defined = [m["mean"] for m in metrics if m["mean"] is not None]
        section = rounded(sum(Decimal(str(v)) for v in defined) / len(defined)) if defined else None
        return {"sample_n": len(selected), "metrics": metrics, "smart_score": section}
    finally:
        workbook.close()


protected = {SUMMARY, *DETAIL.glob("*.xlsx")}
for label in (*LABELS, "Q3"):
    protected.update((RAW / label).glob("*.xlsx"))
protected.update(ROOT / p for p in (
    "src/hangbo/pipeline/cumulative.py", "src/hangbo/pipeline/runtime.py",
    "src/hangbo/summary/table.py", "src/hangbo/survey/stats.py",
))
before = {p: digest(p) for p in protected}
for label, files in RUN_LOG["source_hashes"].items():
    for filename, expected in files.items():
        assert digest(RAW / label / filename) == expected, (label, filename, "source changed since pipeline run")

periods = tuple(calculate_period(
    RAW / label, label, sheet_name=CONFIG["sheet_name"],
    calculation_mode=RUN_LOG["calculation_mode"],
) for label in LABELS)
result = build_cumulative_result(periods)
evidence = []
for period in periods:
    independent = raw_smart_values(period.label)
    stats = period.reports["会展服务商"]
    smart = next(s for s in stats.sections if s.name == "智慧场馆").satisfaction
    assert independent["sample_n"] == stats.matched_row_count
    if smart is None:
        assert independent["smart_score"] is None
    else:
        close(smart, independent["smart_score"])
    counts = {c: sum(bool(r.display_name) and r.values[c] is not None for r in period.rows)
              for c in ("总分", "智慧场馆/服务")}
    evidence.append({"period": period.label, **independent, "valid_customer_type_counts": counts})

valid = [e for e in evidence if e["smart_score"] is not None]
summary_expected = sum(e["smart_score"] / e["valid_customer_type_counts"]["智慧场馆/服务"] for e in valid) / sum(
    1 / e["valid_customer_type_counts"]["智慧场馆/服务"] for e in valid)
detail_expected = sum(e["smart_score"] / e["valid_customer_type_counts"]["总分"] for e in valid) / sum(
    1 / e["valid_customer_type_counts"]["总分"] for e in valid)
summary_row = next(r for r in result.rows if r.display_name == "会展服务商")
close(summary_expected, summary_row.values["智慧场馆/服务"])
aggregated, _ = _aggregate_stats_for_rule(result, "会展服务商", "会展服务商")
close(detail_expected, next(s.satisfaction for s in aggregated.sections if s.name == "智慧场馆"))

workbook = load_workbook(SUMMARY, read_only=True, data_only=True)
try:
    sheet = workbook["汇总表"]
    close(summary_expected, sheet["G6"].value)
    saved_summary = {sheet.cell(i, 2).value: {c: sheet.cell(i, j).value for j, c in enumerate(SUMMARY_COLUMNS, 3)}
                     for i in range(3, 22) if sheet.cell(i, 2).value}
finally:
    workbook.close()
workbook = load_workbook(DETAIL / "会展服务商.xlsx", read_only=True, data_only=True)
try:
    close(detail_expected, workbook["会展服务商"]["B31"].value)
finally:
    workbook.close()

rebuilt_rows = build_summary_rows(load_report_snapshots(DETAIL))
differences = []
for row in rebuilt_rows:
    if row.display_name not in saved_summary:
        continue
    for column in SUMMARY_COLUMNS:
        actual = saved_summary[row.display_name][column]
        from_details = row.values[column]
        if isinstance(actual, (int, float)) and from_details is not None and abs(rounded(actual) - rounded(from_details)) > 1e-9:
            differences.append({"customer_type": row.display_name, "indicator": column,
                                "summary": actual, "rebuilt_from_detail": from_details})

pooled = raw_smart_values("Q3")
pooled_period = calculate_period(RAW / "Q3", "合并样本直接统计", sheet_name=CONFIG["sheet_name"],
                                 calculation_mode=RUN_LOG["calculation_mode"])
close(pooled["smart_score"], next(s.satisfaction for s in pooled_period.reports["会展服务商"].sections if s.name == "智慧场馆"))
assert pooled["sample_n"] == sum(e["sample_n"] for e in evidence)
for path, original in before.items():
    assert digest(path) == original, (path, "audit changed a protected file")

output = ROOT / "logs/verification/2026-Q3-smart-venue-consistency.json"
output.parent.mkdir(parents=True, exist_ok=True)
payload = {"pipeline_source": RUN_LOG["source"], "periods": evidence,
           "summary_G6": summary_expected, "detail_B31": detail_expected,
           "pooled_raw_Q3": pooled, "display_level_mismatches": differences,
           "protected_files_unchanged": len(before), "checks": "all reproduction and integrity assertions passed"}
output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n")
print(json.dumps({"summary_G6": summary_expected, "detail_B31": detail_expected,
                  "pooled_Q3": pooled["smart_score"], "periods": evidence,
                  "display_level_mismatches": len(differences), "protected_files_unchanged": len(before),
                  "audit_log": str(output)}, ensure_ascii=False, indent=2))
