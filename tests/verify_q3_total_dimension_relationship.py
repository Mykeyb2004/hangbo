"""Historical reproduction for d51452c and its pre-fix Q3 workbooks.

This intentionally asserts the old mismatch. For the repaired algorithm use
verify_hierarchical_summary_regression.py instead.
"""

from decimal import Decimal, ROUND_HALF_UP
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

from hangbo.pipeline.cumulative import calculate_period, build_cumulative_result, _aggregate_stats_for_rule
from hangbo.summary.table import load_report_snapshots


ROOT = Path(__file__).resolve().parents[1]


def rounded_mean(values):
    defined = [Decimal(str(value)) for value in values if value is not None]
    return float((sum(defined) / len(defined)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)) if defined else None


if __name__ == "__main__":
    protected = {p for label in ("7月", "8月", "9月", "Q3")
                 for category in ("raw", "satisfaction_detail", "satisfaction_summary")
                 for p in (ROOT / "data" / category / "2026" / label).glob("*.xlsx")}
    before = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in protected}
    periods = tuple(calculate_period(ROOT / "data/raw/2026" / label, label,
                                    sheet_name="问卷数据", calculation_mode="template")
                    for label in ("7月", "8月", "9月"))
    result = build_cumulative_result(periods)
    monthly = []
    weights = []
    for period in periods:
        stats = period.reports["展览主承办"]
        monthly.append({"month": period.label, "total": stats.satisfaction,
                        "sections": {s.name: s.satisfaction for s in stats.sections},
                        "dimension_mean": rounded_mean([s.satisfaction for s in stats.sections])})
        weights.append(result.period_row_weights[(period.label, "展览活动主（承）办", "总分")])
    cumulative, _ = _aggregate_stats_for_rule(result, "展览主承办", "展览活动主（承）办")
    normalized_weights = [weight / sum(weights) for weight in weights]
    reproduced_total = sum(m["total"] * weight for m, weight in zip(monthly, normalized_weights))
    assert abs(reproduced_total - cumulative.satisfaction) < 1e-12
    assert all(m["total"] == m["dimension_mean"] for m in monthly)
    target = next(r for r in load_report_snapshots(ROOT / "data/satisfaction_detail/2026/Q3")
                  if r.source_path.stem == "展览主承办")
    assert target.total == 9.69
    assert rounded_mean(list(target.sections.values())) == 9.76
    workbook = load_workbook(ROOT / "data/satisfaction_summary/2026/Q3/Q3客户类型满意度汇总表.xlsx", data_only=True)
    try:
        row = [workbook["汇总表"].cell(3, col).value for col in range(3, 8)]
        assert row == [9.69, 9.77, 9.55, 9.7, 10]
    finally:
        workbook.close()
    mismatches = []
    for snapshot in load_report_snapshots(ROOT / "data/satisfaction_detail/2026/Q3"):
        mean = rounded_mean(list(snapshot.sections.values()))
        if snapshot.total != mean:
            mismatches.append({"role": snapshot.role_name, "total": snapshot.total,
                               "mean_of_first_level_sections": mean})
    assert all(hashlib.sha256(p.read_bytes()).hexdigest() == digest for p, digest in before.items())
    output = ROOT / "logs/verification/2026-Q3-total-dimension-relationship.json"
    output.parent.mkdir(parents=True, exist_ok=True)
    payload = {"monthly": monthly, "normalized_month_weights": normalized_weights,
               "quarter_unrounded_total": cumulative.satisfaction,
               "quarter_unrounded_sections": {s.name: s.satisfaction for s in cumulative.sections},
               "quarter_published_total": target.total, "quarter_published_sections": target.sections,
               "quarter_dimension_mean_rounded": rounded_mean(list(target.sections.values())),
               "quarter_detail_total_mismatches": mismatches, "protected_excel_unchanged": len(before)}
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(payload, ensure_ascii=False, indent=2))
