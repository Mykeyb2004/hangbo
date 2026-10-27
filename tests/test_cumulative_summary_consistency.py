"""Integration coverage of cumulative workbooks using the local survey fixtures."""

import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import load_workbook

from hangbo.pipeline.cumulative import run_cumulative_batch
from hangbo.summary.table import SUMMARY_COLUMNS, build_summary_rows, load_report_snapshots


ROOT = Path(__file__).resolve().parents[1]


@unittest.skipUnless((ROOT / "data/raw/2026/9月/会展服务商.xlsx").is_file(), "local 2026 survey data required")
class CumulativeSummaryConsistencyTest(unittest.TestCase):
    def test_quarter_summary_uses_detail_scores_including_missing_indicators(self):
        with TemporaryDirectory() as directory:
            output = Path(directory)
            detail_dir = output / "details"
            summary_path = output / "summary.xlsx"
            result = run_cumulative_batch(
                year="2026", batch="Q3", data_root=ROOT / "data",
                detail_output_dir=detail_dir, summary_output_path=summary_path,
                sheet_name="问卷数据", calculation_mode="template",
                january_february_mode="auto", log_path=output / "cumulative.json",
            )
            workbook = load_workbook(summary_path, data_only=True)
            try:
                summary = workbook["汇总表"]
                # Worked example: detail smart-venue score is 9.84, old summary was 9.725.
                self.assertEqual(summary["G6"].value, 9.84)
                detail_rows = build_summary_rows(load_report_snapshots(detail_dir))
                self.assertEqual(result.rows, detail_rows)
                for row_index, row in enumerate(detail_rows, 3):
                    for column_index, indicator in enumerate(SUMMARY_COLUMNS, 3):
                        actual = summary.cell(row_index, column_index).value
                        expected = row.values[indicator]
                        with self.subTest(customer=row.display_name, indicator=indicator):
                            self.assertEqual(actual if isinstance(actual, (int, float)) else None, expected)
                self.assertEqual(summary["C22"].value, result.total_values["总分"])
                self.assertEqual(workbook["周期计算"]["B5"].value, result.total_values["总分"])
                self.assertIn("分表", summary["A24"].value)
                log = json.loads((output / "cumulative.json").read_text())
                self.assertEqual(log["summary_source"], "generated_detail_reports")
                self.assertEqual(log["total"], result.total_values["总分"])
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
